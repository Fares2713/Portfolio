#!/usr/bin/env python3
"""Verify the shipped workbook with genuine LibreOffice formula recalculation."""
from pathlib import Path
from datetime import date,datetime
import shutil,subprocess,tempfile,warnings
from openpyxl import load_workbook
warnings.filterwarnings('ignore',category=UserWarning,module='openpyxl')
root=Path(__file__).resolve().parents[1]
source=root/'assets/rencontres-enjeu/07-classeur-pilotage.xlsx'
w=load_workbook(source)
assert len([s for s in w if s.sheet_state=='visible'])==20
assert len(w['Rétroplanning']['A8:A38'])==31
assert sum(len(s._charts) for s in w)==2
assert w['Paramètres']['B10'].value==80
assert all(w['Paramètres'][f'B{r}'].value is None for r in range(15,19))
assert all(w['Dépenses'].cell(r,c).value is None for r in range(8,48) for c in range(1,11))
assert 'Statuts' in w.defined_names and 'TaskIDs' in w.defined_names
assert len(w['Gantt'].conditional_formatting)==1
engine=shutil.which('soffice') or shutil.which('libreoffice')
assert engine,'LibreOffice is required to verify formula results.'
with tempfile.TemporaryDirectory(prefix='rencontres-xlsx-') as tmp:
 tmp=Path(tmp);out=tmp/'calculated';out.mkdir()
 shutil.copyfile(source,tmp/'baseline.xlsx')
 s=load_workbook(source);p=s['Paramètres'];p['B10']=100;p['B8']=date(2027,4,15);p['B17']=90;p['B18']=81;p['B15']=81;p['B16']=54
 s['Financements']['E8']=1500;s['Dépenses']['A8']='B01';s['Dépenses']['E8']=600;s['Dépenses']['G8']=300;s.save(tmp/'scenario.xlsx')
 s=load_workbook(source)
 for r in range(15,19):s['Paramètres'][f'B{r}']=0
 s['Dépenses']['A8']='ID-INCONNU';s['Dépenses']['E8']=0;s['Dépenses']['G8']=0;s['Paramètres']['B9']=date(2027,4,1);s.save(tmp/'zero.xlsx')
 subprocess.run([engine,f'-env:UserInstallation={tmp.as_uri()}/profile','--headless','--convert-to','xlsx','--outdir',str(out),*[str(tmp/n) for n in ['baseline.xlsx','scenario.xlsx','zero.xlsx']]],check=True,timeout=60,capture_output=True)
 data={n:load_workbook(out/(n+'.xlsx'),data_only=True) for n in ['baseline','scenario','zero']}
 for name,book in data.items():
  errors=[(s.title,c.coordinate,c.value) for s in book for row in s for c in row if c.data_type=='e'];assert not errors,(name,errors)
 b=data['baseline'];assert [b['Budget'][c].value for c in ['H33','H35','H36','H38']]==[2100,210,2310,190]
 assert b['Financements']['E15'].value==2310
 assert b['Tableau de bord']['B13'].value==88
 assert all(b['Indicateurs'][c].value is None for c in ['B9','B10','B11','B12','B13','B14','B16','B17','B18','B21'])
 assert all(b['Rétroplanning'][f'O{r}'].value in ['Cohérent','Sans dépendance'] for r in range(8,39))
 s=data['scenario'];assert s['Budget']['H36'].value==2387
 assert s['Financements']['E15'].value==887
 assert abs(s['Indicateurs']['B11'].value-.9)<1e-10
 assert abs(s['Indicateurs']['B14'].value-2/3)<1e-10
 assert s['Indicateurs']['B21'].value==300
 assert s['Rétroplanning']['H8'].value==datetime(2027,2,18)
 assert s['Rétroplanning']['H20'].value==datetime(2026,10,13),'Completed design date must not move with the event.'
 assert s['Tableau de bord']['C13'].value=='Capacité minimale à augmenter'
 z=data['zero'];assert z['Indicateurs']['B9'].value==0
 assert z['Indicateurs']['B11'].value is None and z['Indicateurs']['B14'].value is None
 assert z['Indicateurs']['B17'].value==0 and z['Indicateurs']['B18'].value==-2310
 assert z['Dépenses']['K8'].value=='ID budget inconnu'
 assert z['Tableau de bord']['B16'].value>0,'Past-due alert should respond to the review date.'
 print('PASS: 20 sheets, baseline budget, blank actuals, dependencies, date/jauge changes, finance, zero denominators, ID and overdue alerts; no formula errors.')
