#!/usr/bin/env python3
"""Build the planning portfolio pack. Dependencies: reportlab, python-docx, openpyxl."""
from pathlib import Path
from datetime import date, timedelta
from decimal import Decimal
from html import escape
import json, zipfile
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, A3, A5
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table as XLTable, TableStyleInfo
from openpyxl.workbook.properties import CalcProperties

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'assets/rencontres-enjeu'
OUT.mkdir(parents=True,exist_ok=True)
INK='#202720'; GOLD='#92681d'; PAPER='#f7f4ec'; MUTED='#536056'
fonts=Path('/usr/share/fonts/truetype/dejavu')
for name,file in [('Body','DejaVuSans.ttf'),('Bold','DejaVuSans-Bold.ttf'),('Editorial','DejaVuSerif.ttf')]:
 pdfmetrics.registerFont(TTFont(name,str(fonts/file)))
pdfmetrics.registerFontFamily('Body',normal='Body',bold='Bold',italic='Body',boldItalic='Bold')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='PBody',fontName='Body',fontSize=9.4,leading=15,textColor=colors.HexColor(INK),spaceAfter=8))
styles.add(ParagraphStyle(name='PHeading',fontName='Editorial',fontSize=22,leading=29,textColor=colors.HexColor(INK),spaceAfter=14))
styles.add(ParagraphStyle(name='PSub',fontName='Bold',fontSize=11,leading=16,textColor=colors.HexColor(GOLD),spaceBefore=12,spaceAfter=8))
styles.add(ParagraphStyle(name='PCell',fontName='Body',fontSize=8,leading=11,textColor=colors.HexColor(INK)))
styles.add(ParagraphStyle(name='PHeader',fontName='Bold',fontSize=8,leading=11,textColor=colors.white))
styles.add(ParagraphStyle(name='PMeta',fontName='Body',fontSize=8,leading=12,textColor=colors.HexColor(MUTED),spaceAfter=12))

def money(v):return f'{v:,.2f}'.replace(',', ' ').replace('.', ',')+' €'
def text(s,style='PBody'):return Paragraph(escape(str(s)).replace('\n','<br/>'),styles[style])
def block(kind,value):return {'kind':kind,'value':value}
def para(s):return block('p',s)
def bullets(*s):return block('bullets',list(s))
def table(headers,rows):return block('table',{'headers':headers,'rows':rows})

E={'title':'Les Rencontres de L’Enjeu','theme':'Comprendre l’information à l’heure des réseaux sociaux et de l’IA',
 'status':'Projet personnel de conception · Événement en préparation', 'period':'Avril 2027',
 'date_reference':'2027-04-08','date_reference_status':'Jeudi 8 avril 2027 : référence de travail, non confirmée',
 'place':'Paris ou Île-de-France · Lieu à confirmer', 'capacity_guests':80,'capacity_room':100,'team_max':8,
 'public_hours':'18 h 30 – 21 h 30','owner':'Farès Benamar','email':'contact.pro.faresb@gmail.com',
 'version':'1.0','prepared':'13 octobre 2026','entry':'Gratuit, sur inscription après confirmation du lieu',
 'financial_status':'Hypothèses TTC à confirmer par devis ; aucun financement acquis'}
COSTS=[('Location de salle',1,600,'Hypothèse ; capacité 100 personnes et accessibilité à confirmer'),
 ('Sonorisation et régie',1,450,'Deux micros, diffusion sonore, technicien selon devis'),
 ('Interventions et modération',1,300,'Enveloppe globale indicative ; conditions à convenir'),
 ('Collation légère',80,3,'Hypothèse par participant ; eau et boissons sans alcool'),
 ('Impression et signalétique',1,90,'Affiches, programme et repères d’accueil'),
 ('Assurance et frais administratifs',1,120,'Hypothèse à préciser avec la structure porteuse'),
 ('Accessibilité et adaptations',1,120,'Provision ; besoins à recueillir avant l’événement'),
 ('Photographie / documentation',1,180,'Option selon devis et consentement des personnes')]
BASE=sum(q*p for _,q,p,_ in COSTS); RESERVE=round(BASE*.10,2); TOTAL=BASE+RESERVE
assert BASE==2100 and TOTAL==2310
TASKS=[
 (-56,-49,'Cadrage','Valider le concept, le public et l’enveloppe','Responsable projet','Dossier de conception','À valider'),
 (-49,-42,'Lieu','Comparer les lieux, les accès et les devis','Logistique','Fiche lieu et devis','À lancer'),
 (-42,-35,'Programme','Identifier puis confirmer les deux intervenants et la modération','Éditorial','Accords écrits et biographies','À lancer'),
 (-42,-28,'Technique','Confirmer salle, assurance, besoins audio et accessibilité','Logistique / régie','Fiche technique','À lancer'),
 (-35,-28,'Supports','Adapter affiche, invitation et page aux informations confirmées','Communication','Supports validés','À préparer'),
 (-28,-21,'Inscriptions','Ouvrir le formulaire et annoncer le rendez-vous','Communication / accueil','Page et formulaire validés','Non ouvert'),
 (-21,-14,'Relais','Diffuser les messages et solliciter les relais pertinents','Communication','Suivi des canaux','À préparer'),
 (-14,-7,'Conduite','Brief des intervenants et validation du déroulé','Responsable projet','Conducteur partagé','À préparer'),
 (-7,-2,'Accueil','Vérifier les inscrits, les besoins et les confirmations','Accueil','Liste d’accueil à accès limité','À préparer'),
 (-2,-1,'Répétition','Tester matériel et consignes avec l’équipe','Régie / logistique','Check-list technique','À préparer'),
 (0,0,'Événement','Déployer le conducteur, accueillir et documenter','Équipe complète','Journal des écarts','Non réalisé'),
 (1,3,'Retour','Envoyer remerciement, questionnaire et ressources','Communication','Message et questionnaire','À préparer'),
 (3,7,'Bilan','Rapprocher dépenses, fréquentation et retours','Responsable projet','Bilan factuel et suites','Non mesuré')]
RUN=[
 ('16:00','16:45','Installation','Logistique / régie','Salle, flux, sonorisation, signalétique','Vérifier accès, issues et capacité avec le lieu'),
 ('16:45','17:15','Tests techniques','Régie','Micros, niveaux, présentation','Tester le plan de secours sans projection'),
 ('17:15','17:45','Brief équipe','Responsable projet','Rôles, accueil, incidents, consentement','Rappeler les responsables du lieu et les contacts'),
 ('17:45','18:15','Brief intervenants','Éditorial / modération','Fil conducteur et timing','Confirmer les noms seulement après accord'),
 ('18:15','18:30','Mise en place accueil','Accueil','Liste, badges et supports','Installer une file accessible et éviter l’affichage des coordonnées'),
 ('18:30','19:00','Accueil du public','Accueil / logistique','Orientation et installation','Noter le nombre d’entrées, sans double comptage'),
 ('19:00','19:10','Ouverture','Responsable projet / modération','Intention et règles d’échange','Présenter les intervenants confirmés'),
 ('19:10','19:45','Table ronde','Modération + deux intervenants','Réseaux sociaux, sources et IA','Limiter les prises de parole et distinguer faits et opinions'),
 ('19:45','20:10','Questions du public','Modération','Questions regroupées par thème','Faire circuler le micro de façon inclusive'),
 ('20:10','20:25','Pause','Logistique','Collation et échanges','Réduire déchets et congestion'),
 ('20:25','20:55','Atelier collectif','Modération / éditorial','Examiner une information et ses sources','Utiliser un exemple documenté, sans exposer de données personnelles'),
 ('20:55','21:05','Synthèse','Modération','Trois repères pratiques et ressources','Indiquer comment retrouver les ressources'),
 ('21:05','21:30','Échanges informels','Équipe complète','Rencontres et retours','Veiller à l’horaire et aux participants ayant besoin d’aide'),
 ('21:30','21:45','Sortie et remise en état','Logistique / régie','Rangement et restitution','Vérifier salle, matériel et objets trouvés')]
RISKS=[
 ('Lieu ou financement non confirmé','Fort','Fort','Ne pas ouvrir les inscriptions avant décision écrite ; comparer au moins deux options','Responsable projet','Non vérifié'),
 ('Désistement d’un intervenant','Moyen','Fort','Confirmer par écrit ; prévoir un format à un intervenant et une séquence de questions étendue','Éditorial','À préparer'),
 ('Inscrits absents','Moyen','Moyen','Rappel à J-2, confirmation simple et liste d’attente sans dépasser la jauge','Accueil','À préparer'),
 ('Panne audio / projection','Moyen','Fort','Test J-1 et le jour même ; conducteur et supports imprimés de secours','Régie','À tester'),
 ('Besoin d’accessibilité non anticipé','Moyen','Fort','Vérifier le lieu et offrir un contact pour recueillir les besoins, sans demander de diagnostic médical','Logistique','À vérifier'),
 ('Photo ou donnée utilisée sans accord','Moyen','Fort','Informer, recueillir l’accord pour les usages prévus et limiter les accès à la liste d’inscrits','Communication','À préparer'),
 ('Incident dans le lieu','Faible','Fort','Suivre le protocole du lieu ; repérer issues, personnel référent et moyens d’alerte','Logistique','À confirmer avec le lieu'),
 ('Dérive budgétaire','Moyen','Fort','Aucun engagement sans validation ; réserve de 10 % et arbitrage des options avant commande','Responsable projet','À suivre')]
RACI=[
 ('Concept et budget','A/R','C','C','C','I'),('Programme et intervenants','A','R','C','I','I'),
 ('Lieu et technique','A','C','I','R','C'),('Supports et diffusion','A','C','R','C','I'),
 ('Inscriptions et accueil','A','I','C','C','R'),('Déroulé et brief équipe','A/R','C','C','C','C'),
 ('Bilan et suivi','A','C','R','C','C')]

DOCS=[
('01-dossier-conception','Dossier de conception',[
 ('Le projet en une page',[
  para('Une rencontre publique pour comprendre comment circulent les informations, comment les vérifier et ce que l’IA change dans nos pratiques. Le format associe discussion, questions et mise en pratique, dans l’esprit du média indépendant associatif L’Enjeu.'),
  table(['Repère','Hypothèse de conception'],[['Statut',E['status']],['Période',E['period']+' ; date exacte à confirmer'],['Lieu',E['place']],['Public cible','Lecteurs du média, étudiants, acteurs associatifs et personnes intéressées par l’information'],['Jauge','80 participants visés ; salle de 100 places minimum, équipe jusqu’à 8 personnes'],['Format','Une soirée, deux intervenants proposés et une modération ; aucun intervenant confirmé'],['Accès',E['entry']],['Budget','2 310 € TTC prévisionnels ; plafond de travail 2 500 €']]),
  para('Ce dossier est une production de conception. Aucun événement, partenariat, réservation ou résultat de fréquentation n’est présenté comme acquis.')]),
 ('Objectifs et choix éditoriaux',[
  bullets('Rendre accessibles les enjeux de fiabilité de l’information, sans ton prescriptif ni jargon.','Créer un échange direct entre le média et son public.','Donner trois repères pratiques : identifier la source, chercher la preuve, vérifier le contexte.','Expérimenter un format reproductible pour une éventuelle deuxième édition.'),
  table(['Objectif de travail','Indicateur prévu','Mode de mesure'],[['Viser 80 participants','Entrées réelles, distinctes des inscriptions','Comptage à l’accueil ; résultat à renseigner'],['Faciliter une expérience utile','Appréciation de l’utilité et commentaires','Questionnaire volontaire ; aucun score actuel'],['Garder un budget maîtrisé','Dépenses et écart au prévisionnel','Factures et engagements validés'],['Documenter les suites','Ressources publiées et pistes retenues','Bilan à J+7']]),
  para('Les cibles sont des intentions de conception. Les résultats seront renseignés uniquement après réalisation et à partir de données effectivement collectées.')]),
 ('Format et programme proposé',[
  para('Thème : '+E['theme']+'. Format : une table ronde de 35 minutes, 25 minutes de questions, une pause puis un atelier collectif. Deux profils sont à rechercher : une personne travaillant sur les médias / l’éducation à l’information et une personne travaillant sur les usages de l’IA. La modération reste à désigner.'),
  table(['Horaire','Séquence'],[[a+' – '+b,c] for a,b,c,*_ in RUN if a>='18:30' and a<='21:05']),
  para('Questions de préparation : pourquoi une information attire-t-elle notre attention ? Quelles vérifications sont possibles avant de partager ? Comment contextualiser un contenu produit ou amplifié par l’IA ? Les exemples utilisés devront être sourcés et validés avant présentation.')]),
 ('Publics, lieu et accessibilité',[
  bullets('Lieu à rechercher à Paris ou en Île-de-France, accessible en transports, disposant d’une salle adaptée à au moins 100 personnes.','Prévoir un accueil lisible, des places assises et un cheminement accessible ; vérifier les équipements avec le gestionnaire.','Publier les accès, horaires et conditions d’accueil seulement après confirmation.','Prévoir un contact pour les besoins particuliers, sans collecte de renseignements médicaux.','Jauge d’inscription envisagée : 80 personnes ; équipe maximale de 8 personnes, soit 88 occupants prévus avant vérification du lieu.'),
  table(['Critère de sélection du lieu','À vérifier'],[['Capacité et disposition','Jauge autorisée, visibilité, circulation et espaces de pause'],['Accès','Transports, accès PMR et toilettes accessibles'],['Technique','Micros, prises, acoustique et appui technique'],['Conditions','Devis, assurance demandée, horaires et restitution']])]),
 ('Budget et financement envisagés',[
  table(['Poste','Quantité','Prix TTC','Total TTC'],[[label,str(q),money(u),money(q*u)] for label,q,u,_ in COSTS]),
  table(['Synthèse','Montant'],[['Sous-total',money(BASE)],['Réserve de 10 %',money(RESERVE)],['Total prévisionnel',money(TOTAL)],['Plafond de travail','2 500,00 €'],['Marge par rapport au plafond','190,00 €']]),
  para('Scénario de financement : contribution envisagée de 1 500 € par la structure porteuse, à valider ; solde de 810 € à rechercher ou à couvrir par une réduction de dépenses. Aucun financement ni soutien extérieur n’est confirmé. Tous les montants doivent être rapprochés de devis réels avant engagement.')]),
 ('Décisions avant lancement',[
  table(['Décision','Responsable proposé','Preuve attendue'],[['Valider le portage et le budget','Responsable projet / structure porteuse','Décision écrite et capacité de financement'],['Confirmer date, lieu et assurance','Responsable projet / logistique','Accords et conditions du lieu'],['Confirmer les intervenants','Éditorial','Accord écrit et programme validé'],['Autoriser la diffusion et l’inscription','Responsable projet','Informations confirmées et modalités de données précisées']]),
  para('Farès Benamar est identifié comme concepteur du dossier. Les rôles opérationnels sont proposés et les membres de l’équipe restent à désigner. La réalisation de l’événement dépend de ces validations.'),
  para('Contact de conception : '+E['email']+' · Portfolio : portfolio-fares-benamar.vercel.app')])]),
('02-dossier-operationnel','Organisation & conduite',[
 ('Organisation et responsabilités',[
  para('Un noyau de coordination organise le projet ; une équipe de huit personnes maximum est envisagée le jour de l’événement. Les rôles ci-dessous ne désignent pas une équipe déjà constituée.'),
  table(['Rôle','Mission proposée'],[['Responsable projet','Cadrage, budget, décisions, suivi et brief général'],['Éditorial / modération','Fil conducteur, intervenants, questions et rythme'],['Communication','Supports, diffusion, ressources et retours'],['Logistique / régie','Lieu, technique, installation et restitution'],['Accueil','Inscriptions, orientation et comptage']]),
  table(['Action','Projet','Éditorial','Communication','Logistique','Accueil'],RACI),
  para('R = réalise ; A = assume la décision ; C = consulté ; I = informé. Les personnes et prestataires sont à confirmer avant lancement.')]),
 ('Rétroplanning',[
  para('Le jeudi 8 avril 2027 sert uniquement de référence de travail au calendrier. La communication publique indique avril 2027, sans annoncer un jour confirmé. Les échéances relatives devront être recalculées après validation de la date.'),
  table(['Échéance','Phase','Livrable / décision'],[[f'J{a:+d} à J{b:+d}',phase,deliverable] for a,b,phase,action,owner,deliverable,status in TASKS]),
  para('Point d’équipe hebdomadaire de 20 minutes proposé : décisions, actions de la semaine, blocages, écarts budget et prochaine échéance. Le classeur joint permet d’attribuer les tâches et de renseigner leur avancement.')]),
 ('Conducteur du jour',[
  table(['Début','Fin','Séquence','Pilote proposé'],[[a,b,c,d] for a,b,c,d,*_ in RUN]),
  para('La fin publique est prévue à 21 h 30 et la restitution à 21 h 45, sous réserve des horaires du lieu. Le détail technique et les vigilances sont disponibles dans la feuille « Déroulé » du classeur.')]),
 ('Fiche technique et accueil',[
  bullets('Prévoir deux micros pour les intervenants et une solution de prise de parole pour le public, à confirmer avec la régie.','Tester la sonorisation et la projection ; conserver une conduite possible sans écran.','Installer un point d’accueil avec programme, repères et contact de l’équipe.','Prévoir une liste d’accueil à accès limité, un comptage des entrées et une procédure simple pour la liste d’attente.','Recueillir les besoins d’accès avant la soirée et vérifier les adaptations avec le lieu.','Repérer les issues et les consignes du gestionnaire ; désigner un référent logistique.','Les prises de vue sont optionnelles. Informer le public, respecter les personnes ne souhaitant pas être photographiées et obtenir les accords nécessaires aux usages annoncés.'),
  table(['Moment','Contrôle indispensable'],[['Avant ouverture','Capacité, accès, matériel, horaires, équipe et programme confirmés'],['Pendant','Comptage, rythme, circulation, incidents et besoins particuliers'],['Après','Restitution du lieu, inventaire, dépenses et retour d’équipe']])]),
 ('Risques et solutions de repli',[
  table(['Risque','Prévention / repli','Pilote'],[[a,d,e] for a,b,c,d,e,f in RISKS]),
  para('La liste est un outil de préparation, pas une attestation de conformité du lieu. Les mesures opérationnelles doivent être validées avec le gestionnaire et les personnes compétentes.')]),
 ('Données, suivi et bilan',[
  bullets('Le formulaire d’inscription n’est pas encore ouvert. Ne recueillir que les informations nécessaires : nom, e-mail, nombre de places et besoin d’accès formulé librement si la personne le souhaite.','L’inscription à une newsletter doit être une démarche distincte et facultative.','Définir avant ouverture le responsable du traitement, la notice d’information, la durée de conservation et les accès. Proposition de travail : suppression de la liste d’accueil après le bilan, au plus tard J+30, hors obligations justifiées.','Questionnaire de retour volontaire : utilité, clarté, accueil, point à améliorer et souhait de revenir.','Bilan J+7 : entrées réelles, inscriptions, dépenses engagées, retours et trois améliorations prioritaires. Aucune donnée de fréquentation n’est actuellement renseignée.'),
  table(['Question du bilan','Source'],[['Qu’a-t-on effectivement réalisé ?','Conducteur et journal des écarts'],['Qui a participé ?','Comptage des entrées et inscriptions, sans publier de données nominatives'],['Quel coût ?','Devis, engagements et factures'],['Qu’est-ce qui doit évoluer ?','Retour d’équipe et questionnaire volontaire']])])]),
('03-kit-communication','Kit de communication',[
 ('Positionnement et règles de diffusion',[
  para('Promesse : une soirée pour prendre du recul sur l’information et repartir avec des repères pratiques. Ton : accessible, précis, ouvert à la discussion. Éviter les promesses de vérité absolue, les témoignages fictifs et les noms d’intervenants non confirmés.'),
  table(['Information','Version de conception'],[['Nom',E['title']],['Thème',E['theme']],['Période',E['period']+' · date exacte à confirmer'],['Zone',E['place']],['Accès',E['entry']],['Statut','Projet en préparation ; inscriptions non ouvertes']]),
  para('Les textes suivants sont prêts à adapter. Une annonce ferme nécessite la confirmation du jour, de l’adresse, du programme, des intervenants et du dispositif d’inscription. Aucun lien d’inscription fictif n’est utilisé.')]),
 ('Invitation par e-mail',[
  para('Objet : Les Rencontres de L’Enjeu, une première édition envisagée en avril 2027'),
  para('Bonjour,\n\nComment prendre du recul sur une information qui circule vite ? Que changent les réseaux sociaux et l’IA dans notre façon de nous informer ?\n\nNous préparons une première édition des Rencontres de L’Enjeu en avril 2027, à Paris ou en Île-de-France. Au programme envisagé : une table ronde, vos questions et un atelier collectif pour identifier les sources et remettre les informations en contexte.\n\nL’accès est prévu gratuit, sur inscription. La date exacte, le lieu et les intervenants seront annoncés après confirmation.\n\nPour être informé de l’ouverture des inscriptions, vous pouvez écrire à '+E['email']+'.\n\nÀ bientôt,\nFarès Benamar\nProjet Les Rencontres de L’Enjeu'),
  para('Ce message est une proposition de texte. Aucun e-mail n’a été envoyé et aucune liste de destinataires n’est fournie.')]),
 ('Publications et relais',[
  para('Publication LinkedIn proposée'),
  para('Nous préparons Les Rencontres de L’Enjeu : un rendez-vous pour discuter de l’information à l’heure des réseaux sociaux et de l’IA.\n\nUne table ronde, des questions et un atelier pour repartir avec des repères concrets : d’où vient une information, sur quelles preuves repose-t-elle, dans quel contexte faut-il la lire ?\n\nPremière édition envisagée en avril 2027, à Paris ou en Île-de-France. Accès prévu gratuit sur inscription. Date, lieu et intervenants à confirmer.\n\nVous souhaitez recevoir les informations lors de l’ouverture ? Contact : '+E['email']+'.'),
  para('Publication courte proposée'),
  para('Comprendre l’information. Questionner les sources. Prendre le temps d’échanger. Les Rencontres de L’Enjeu sont en préparation pour avril 2027. Une soirée gratuite, avec table ronde et atelier. Date et lieu à confirmer.'),
  para('Message de relais proposé'),
  para('Bonjour, nous concevons une rencontre publique autour de l’information, des réseaux sociaux et de l’IA, envisagée en avril 2027. Votre public pourrait être intéressé. Seriez-vous disponible pour un échange sur le format et un éventuel relais une fois les informations confirmées ? Cette prise de contact n’implique aucun partenariat acquis.')]),
 ('Rappel et messages après inscription',[
  para('Modèle de confirmation à utiliser seulement après ouverture effective'),
  para('Objet : Votre inscription aux Rencontres de L’Enjeu\nBonjour, votre inscription est confirmée pour [date confirmée], de 18 h 30 à 21 h 30, à [lieu et adresse confirmés]. Retrouvez le programme et les accès sur [page validée]. Si vous ne pouvez plus venir, prévenez-nous à '+E['email']+'. Pour un besoin d’accès particulier, vous pouvez également nous écrire.'),
  para('Rappel J-2 à adapter'),
  para('Nous avons rendez-vous dans deux jours pour Les Rencontres de L’Enjeu. Accueil dès 18 h 30 ; début des échanges à 19 h. Vérifiez le lieu, l’itinéraire et le programme dans votre confirmation. Un empêchement ? Un message permet de libérer votre place.'),
  para('Remerciement J+1 à utiliser seulement après réalisation'),
  para('Merci d’avoir participé aux Rencontres de L’Enjeu. Nous vous transmettons les ressources citées : [liens vérifiés]. Si vous le souhaitez, partagez votre retour dans [questionnaire actif]. Votre avis nous aidera à améliorer une éventuelle prochaine édition.'),
  para('Les champs entre crochets sont des informations manquantes à compléter ; les modèles ne doivent pas être diffusés ainsi.')]),
 ('Calendrier et indicateurs',[
  table(['Moment','Support','But'],[['Après validation, J-28','Page, invitation et annonce','Présenter une information complète et ouvrir les inscriptions'],['J-21','Publication et relais ciblés','Faire connaître le rendez-vous auprès des publics concernés'],['J-14','Programme et angle éditorial','Expliquer ce que les participants pourront découvrir'],['J-7','Informations pratiques','Réduire les questions d’accès et d’organisation'],['J-2','Rappel aux inscrits','Confirmer la présence et libérer les places inutilisées'],['J+1','Remerciement et ressources','Prolonger les échanges'],['J+7','Bilan factuel','Évaluer les objectifs et décider des suites']]),
  table(['Indicateur prévu','Calcul / précaution'],[['Inscriptions par source','Formulaire : source facultative ou liens de campagne, sans déductions inventées'],['Taux de présence','Entrées uniques / inscriptions confirmées × 100'],['Taux de réponse au questionnaire','Réponses / participants contactés × 100'],['Utilité perçue','Distribution des réponses ; ne pas présenter une cible comme un résultat'],['Écart budgétaire','Dépenses réelles - prévisionnel, avec les mêmes postes']]),
  para('Le classeur contient les champs et les formules ; les cellules de résultats restent vides tant que les données n’existent pas.')]),
 ('Direction graphique et supports',[
  para('Palette : ivoire, encre et or, inspirée de l’identité du média. Typographie éditoriale pour les titres et caractères sans empattement pour les informations pratiques. Maintenir un contraste suffisant et une hiérarchie courte : nom, sujet, période, format, statut.'),
  bullets('Affiche A3 : annonce du concept et de la période envisagée ; aucune fausse adresse ni aucun intervenant fictif.','Invitation A5 : format et modalités prévues, avec contact de conception.','Programme A4 : séquences horaires de la soirée, sans noms non confirmés.','Page de présentation HTML : aperçu du rendez-vous, programme et contact ; aucun formulaire actif ni collecte de données.'),
  para('Les supports sont des productions de conception destinées au portfolio. Adapter les mentions provisoires après validation, sans attribuer une réalisation effective à cet exercice.')])])]

# Store readable source data so future revisions do not require editing generated PDFs.
SOURCE={'event':E,'budget':{'lines':COSTS,'base':BASE,'reserve_rate':.10,'reserve':RESERVE,'total':TOTAL},'tasks':TASKS,'run':RUN,'risks':RISKS,'raci':RACI,'documents':DOCS}
(OUT/'sources-projet.json').write_text(json.dumps(SOURCE,ensure_ascii=False,indent=2)+'\n')

def footer(c,doc):
 c.saveState();w,h=A4
 c.setStrokeColor(colors.HexColor(GOLD));c.setLineWidth(.8);c.line(42,h-34,w-42,h-34)
 c.setFont('Bold',8);c.setFillColor(colors.HexColor(INK));c.drawString(42,h-26,'L’Enjeu · Les Rencontres')
 c.setFont('Body',7);c.setFillColor(colors.HexColor(MUTED));c.drawString(42,25,'Conception · Avril 2027 envisagé · V1.0 · Farès Benamar')
 c.drawRightString(w-42,25,str(doc.page));c.restoreState()

def pdf_document(stem,title,sections):
 doc=SimpleDocTemplate(str(OUT/(stem+'.pdf')),pagesize=A4,rightMargin=42,leftMargin=42,topMargin=52,bottomMargin=45,
   title=E['title']+' | '+title,author=E['owner'],subject=E['status'])
 flow=[]
 for i,(heading,blocks) in enumerate(sections):
  if i:flow.append(PageBreak())
  if i==0:
   flow.extend([text(title,'PHeading'),text(E['title']+' · '+E['theme'],'PSub'),text(E['status']+'\n'+E['period']+' · Date, lieu et intervenants à confirmer','PMeta')])
  flow.append(text(f'{i+1:02} / {heading}','PSub'))
  for b in blocks:
   if b['kind']=='p':flow.append(text(b['value']))
   elif b['kind']=='bullets':
    flow.extend(text('• '+item) for item in b['value'])
   else:
    v=b['value'];cols=len(v['headers']);width=(A4[0]-84)/cols
    rows=[[text(h,'PHeader') for h in v['headers']]]+[[text(cell,'PCell') for cell in row] for row in v['rows']]
    t=Table(rows,colWidths=[width]*cols,repeatRows=1,hAlign='LEFT')
    t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(INK)),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor(PAPER),colors.white]),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#d9ddd4'))]))
    flow.extend([t,Spacer(1,12)])
 doc.build(flow,onFirstPage=footer,onLaterPages=footer)

def docx_document(stem,title,sections):
 d=Document();sec=d.sections[0];sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
 normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10);normal.paragraph_format.space_after=Pt(7)
 for style_name in ['Title','Heading 1','Heading 2']:
  d.styles[style_name].font.name='Georgia'
  d.styles[style_name].font.color.rgb=RGBColor.from_string('202720' if style_name=='Title' else '92681D')
 d.core_properties.title=E['title']+' | '+title;d.core_properties.author=E['owner'];d.core_properties.subject=E['status']
 sec.header.paragraphs[0].text='L’Enjeu · Les Rencontres'
 sec.footer.paragraphs[0].text='Projet de conception · Avril 2027 envisagé · V1.0 · Farès Benamar'
 d.add_heading(title,0);d.add_paragraph(E['title']);d.add_paragraph(E['status']+'\nDate, lieu et intervenants à confirmer')
 for i,(heading,blocks) in enumerate(sections):
  if i:d.add_page_break()
  d.add_heading(f'{i+1:02} / {heading}',1)
  for b in blocks:
   if b['kind']=='p':d.add_paragraph(b['value'])
   elif b['kind']=='bullets':
    for item in b['value']:d.add_paragraph(item,style='List Bullet')
   else:
    v=b['value'];t=d.add_table(rows=1,cols=len(v['headers']));t.style='Light Shading Accent 1'
    for cell,val in zip(t.rows[0].cells,v['headers']):cell.text=str(val)
    for row in v['rows']:
     for cell,val in zip(t.add_row().cells,row):cell.text=str(val)
    for cell in t.rows[0].cells:
     shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'202720');cell._tc.get_or_add_tcPr().append(shading)
     for paragraph in cell.paragraphs:
      for run in paragraph.runs:run.font.bold=True;run.font.color.rgb=RGBColor(255,255,255)
    d.add_paragraph()
 d.save(OUT/(stem+'.docx'))

for stem,title,sections in DOCS:
 pdf_document(stem,title,sections);docx_document(stem,title,sections)

# Operational workbook: inputs stay explicit; unknown results remain blank.
wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
def sheet(name,headers,rows,widths):
 ws=wb.create_sheet(name);ws.append(headers)
 for row in rows:ws.append(list(row))
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.sheet_view.showGridLines=False
 for cell in ws[1]:cell.font=Font(name='Calibri',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='202720');cell.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=32
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name='Calibri',size=11,color='202720');cell.alignment=Alignment(wrap_text=True,vertical='top')
  ws.row_dimensions[row[0].row].height=45
 for i,width in enumerate(widths,1):ws.column_dimensions[__import__('openpyxl').utils.get_column_letter(i)].width=width
 ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
 ws.print_title_rows='1:1'
 return ws
sheet('Mode emploi',['Repère','Instruction'],[
 ['Statut',E['status']],['Période publique',E['period']+' ; date exacte à confirmer'],['Date de travail','08/04/2027 = référence non confirmée, à modifier après validation'],['Cellules bleues','Données ou hypothèses modifiables'],['Formules','Calcul automatique à l’ouverture dans un tableur compatible'],['Données inconnues','Ne pas remplacer les cellules de résultats vides par des résultats fictifs'],['Budget','Tous les montants sont des hypothèses TTC à confirmer par devis'],['Rétroplanning','Dates calculées depuis la cellule B2 de la feuille Paramètres'],['Données personnelles','Ce modèle ne contient aucune liste de participants'],['Version','V1.0 · 13 octobre 2026 · Farès Benamar']],[25,100])
params=sheet('Paramètres',['Paramètre','Valeur'],[['Date de référence',date(2027,4,8)],['Participants visés',80],['Capacité minimale salle',100],['Équipe maximale',8],['Statut','Date de référence non confirmée ; diffusion publique : avril 2027']],[32,70]);params['B2'].number_format='dd/mm/yyyy'
for row in range(2,6):params.cell(row,2).font=Font(color='173EC4')
b=sheet('Budget',['Poste','Quantité','Prix unitaire TTC','Total TTC','Hypothèse / validation'],[(label,q,u,f'=B{i}*C{i}',note) for i,(label,q,u,note) in enumerate(COSTS,2)],[35,12,20,20,85])
b['B5']="='Paramètres'!B3" # Collation quantity tracks the target attendance.
b.append(['Sous-total','','','=SUM(D2:D9)','Hypothèses, pas des dépenses engagées'])
b.append(['Réserve',.10,'','=ROUND(D10*B11,2)','Taux modifiable'])
b.append(['Total prévisionnel','','','=D10+D11','Montants TTC à confirmer'])
b.append(['Plafond de travail','','',2500,'À valider avant engagement'])
b.append(['Marge disponible','','','=D13-D12','Peut devenir négative après modification des hypothèses'])
b.append(['Contribution envisagée','','',1500,'Scénario non validé ; aucun financement acquis'])
b.append(['Solde à financer','','','=MAX(0,D12-D15)','À couvrir ou arbitrer, sans partenaire confirmé'])
for row in b.iter_rows(min_row=2,max_row=16):
 for cell in row:
  cell.alignment=Alignment(wrap_text=True,vertical='top');cell.font=Font(name='Calibri',size=11,color='173EC4' if cell.column in [2,3] and not str(cell.value).startswith('=') else '202720')
 for col in [3,4]:b.cell(row[0].row,col).number_format='#,##0.00 "€"'
b['B11'].number_format='0%';b.auto_filter.ref='A1:E9'
retro=sheet('Rétroplanning',['Début relatif','Fin relative','Phase','Action','Pilote proposé','Livrable','Statut','Début calculé','Fin calculée'],[(a,z,phase,action,owner,deliverable,status,f"='Paramètres'!$B$2+A{i}",f"='Paramètres'!$B$2+B{i}") for i,(a,z,phase,action,owner,deliverable,status) in enumerate(TASKS,2)],[15,15,20,75,30,45,20,20,20])
for row in range(2,len(TASKS)+2):
 for col in [8,9]:retro.cell(row,col).number_format='dd/mm/yyyy'
sheet('Déroulé',['Début','Fin','Séquence','Pilote proposé','Préparation','Vigilance'],RUN,[12,12,28,30,55,85])
sheet('RACI',['Action','Responsable projet','Éditorial','Communication','Logistique','Accueil'],RACI,[40,23,23,23,23,23])
sheet('Risques',['Risque','Probabilité','Impact','Prévention / repli','Pilote','Statut'],RISKS,[40,20,20,100,25,20])
sheet('Diffusion',['Échéance','Canal','Contenu','Objectif','Statut'],[[-28,'Page et e-mail','Programme validé, informations pratiques et inscription','Ouverture après validation','À préparer'],[-21,'Réseaux et relais','Angle éditorial et invitation','Faire connaître le rendez-vous','À préparer'],[-14,'Réseaux','Programme et présentation des intervenants confirmés','Préciser la valeur de la soirée','À préparer'],[-7,'Page et e-mail','Accès, horaires et conditions','Préparer la venue','À préparer'],[-2,'E-mail inscrits','Rappel et désistement','Confirmer / libérer les places','À préparer'],[1,'E-mail participants','Remerciement et ressources','Prolonger les échanges','À préparer'],[7,'Bilan interne','Données vérifiées et améliorations','Décider des suites','Non mesuré']],[15,25,75,45,20])
kpi=sheet('Indicateurs',['Indicateur','Valeur','Statut / définition'],[['Participants visés',"='Paramètres'!B3",'Objectif, pas un résultat'],['Inscriptions confirmées',None,'À renseigner après ouverture'],['Entrées uniques',None,'À renseigner après l’événement'],['Taux de présence','=IF(COUNT(B3:B4)<2,"",IFERROR(B4/B3,""))','Entrées uniques / inscriptions confirmées'],['Questionnaires envoyés',None,'À renseigner après l’événement'],['Réponses reçues',None,'À renseigner après l’événement'],['Taux de réponse','=IF(COUNT(B6:B7)<2,"",IFERROR(B7/B6,""))','Réponses / questionnaires envoyés'],['Dépenses réelles TTC',None,'À renseigner à partir des dépenses vérifiées'],['Budget prévisionnel',"='Budget'!D12",'Hypothèses, 2 310 € au départ'],['Écart réel - prévisionnel','=IF(COUNT(B9)=0,"",B9-B10)','Un écart ne se calcule pas avant saisie des dépenses']],[35,25,90])
for row in [5,8]:kpi.cell(row,2).number_format='0.0%'
for row in [9,10,11]:kpi.cell(row,2).number_format='#,##0.00 "€"'
for row in [3,4,6,7,9]:kpi.cell(row,2).fill=PatternFill('solid',fgColor='E8EEFF');kpi.cell(row,2).font=Font(color='173EC4')
sheet('Décisions',['Sujet','Décision / preuve','Responsable proposé','Échéance','Statut'],[['Portage et financement',None,'Responsable projet','Avant diffusion','Non confirmé'],['Date et lieu',None,'Responsable projet / logistique','J-42','Non confirmé'],['Intervenants',None,'Éditorial','J-35','Non confirmé'],['Assurance et technique',None,'Logistique','J-28','Non confirmé'],['Inscriptions et données',None,'Responsable projet / accueil','Avant ouverture','Non ouvert'],['Dépenses engagées',None,'Responsable projet','Avant commande','Non engagé']],[35,85,35,22,22])
wb.save(OUT/'07-classeur-pilotage.xlsx')

# Professionally typeset public-facing concept pieces, with provisional status on every support.
def small_footer(c,w):
 c.setFillColor(colors.HexColor(MUTED));c.setFont('Body',7);c.drawString(28,20,'Projet de conception · Avril 2027 envisagé · Informations à confirmer')
def poster():
 w,h=A3;c=canvas.Canvas(str(OUT/'04-affiche-a3.pdf'),pagesize=A3);c.setTitle(E['title']+' | Affiche de conception');c.setAuthor(E['owner'])
 c.setFillColor(colors.HexColor(PAPER));c.rect(0,0,w,h,fill=1,stroke=0)
 c.setFillColor(colors.HexColor(GOLD));c.rect(42,h-80,65,8,fill=1,stroke=0)
 c.setFillColor(colors.HexColor(INK));c.setFont('Editorial',40);c.drawString(42,h-135,'L’Enjeu.')
 c.setFont('Body',10);c.drawString(42,h-164,'UN RENDEZ-VOUS POUR COMPRENDRE ET ÉCHANGER')
 for y,t in [(h-280,'Les Rencontres'),(h-339,'de L’Enjeu')]:c.setFont('Editorial',48);c.drawString(42,y,t)
 c.setStrokeColor(colors.HexColor(GOLD));c.line(42,h-380,w-42,h-380)
 c.setFont('Bold',24)
 for y,t in [(h-442,'Comprendre l’information'),(h-478,'à l’heure des réseaux'),(h-514,'sociaux et de l’IA')]:c.drawString(42,y,t)
 c.setFont('Editorial',30);c.drawString(42,h-602,'Avril 2027')
 c.setFont('Body',13)
 for y,t in [(h-638,'Paris / Île-de-France · Lieu à confirmer'),(h-663,'Table ronde · Questions · Atelier collectif'),(h-688,'Accès prévu gratuit, sur inscription')]:c.drawString(42,y,t)
 c.setFillColor(colors.HexColor(GOLD));c.setFont('Bold',11);c.drawString(42,145,'PREMIÈRE ÉDITION EN PRÉPARATION')
 c.setFillColor(colors.HexColor(MUTED));c.setFont('Body',9);c.drawString(42,121,'Date exacte et intervenants à confirmer. Inscriptions non ouvertes.');c.drawString(42,99,E['email'])
 small_footer(c,w);c.save()
def invitation():
 w,h=A5;c=canvas.Canvas(str(OUT/'05-invitation-a5.pdf'),pagesize=A5);c.setTitle(E['title']+' | Invitation de conception');c.setAuthor(E['owner']);c.setFillColor(colors.HexColor(PAPER));c.rect(0,0,w,h,fill=1,stroke=0)
 y=h-40
 for font,size,t,gap in [('Editorial',24,'L’Enjeu.',46),('Bold',9,'ÉDITION PILOTE EN PRÉPARATION',35),('Editorial',25,'Les Rencontres',31),('Editorial',25,'de L’Enjeu',46),('Bold',12,'Comprendre l’information',20),('Body',11,'Réseaux sociaux et IA :',19),('Body',11,'prendre du recul, vérifier, échanger.',38),('Editorial',20,'Avril 2027',28),('Body',10,'Paris / Île-de-France · Lieu à confirmer',24),('Body',10,'18 h 30 – 21 h 30 · Horaires envisagés',22),('Body',10,'Table ronde, questions et atelier collectif',22),('Body',10,'Accès prévu gratuit sur inscription',30),('Body',9,'Date et intervenants à confirmer.',18),('Body',9,'Inscriptions non ouvertes.',23),('Body',8,E['email'],15)]:
  c.setFont(font,size);c.setFillColor(colors.HexColor(INK));c.drawString(28,y,t);y-=gap
 small_footer(c,w);c.save()
def programme():
 pdf_document('06-programme-public','Programme envisagé', [('Une soirée pour prendre du recul',[
  para(E['title']+' · '+E['theme']),para('Avril 2027 · Paris ou Île-de-France. Date, lieu et intervenants à confirmer. Les horaires ci-dessous sont une proposition de conception.'),
  table(['Horaire','Au programme'],[[a+' – '+b,c] for a,b,c,*_ in RUN if a>='18:30' and a<='21:05']),
  para('Accès prévu gratuit sur inscription après confirmation du lieu. Deux intervenants et une modération sont envisagés ; aucun nom n’est annoncé avant accord.'),para('Pour recevoir les informations de préparation : '+E['email'])])])
poster();invitation();programme()

# Vector poster for the portfolio thumbnail and full-size on-screen viewing.
svg='''<svg xmlns="http://www.w3.org/2000/svg" width="840" height="1188" viewBox="0 0 840 1188"><rect width="840" height="1188" fill="#f7f4ec"/><rect x="60" y="62" width="90" height="10" fill="#92681d"/><g fill="#202720"><text x="60" y="150" font-family="Georgia,serif" font-size="62">L’Enjeu.</text><text x="60" y="202" font-family="Arial,sans-serif" font-size="15" letter-spacing="2">COMPRENDRE. QUESTIONNER. ÉCHANGER.</text><text x="60" y="345" font-family="Georgia,serif" font-size="72">Les Rencontres</text><text x="60" y="425" font-family="Georgia,serif" font-size="72">de L’Enjeu</text><path d="M60 475H780" stroke="#92681d"/><text x="60" y="560" font-family="Arial,sans-serif" font-size="34" font-weight="bold">Comprendre l’information</text><text x="60" y="610" font-family="Arial,sans-serif" font-size="34">à l’heure des réseaux</text><text x="60" y="660" font-family="Arial,sans-serif" font-size="34">sociaux et de l’IA</text><text x="60" y="780" font-family="Georgia,serif" font-size="48">Avril 2027</text><text x="60" y="830" font-family="Arial,sans-serif" font-size="21">Paris / Île-de-France · Lieu à confirmer</text><text x="60" y="875" font-family="Arial,sans-serif" font-size="21">Table ronde · Questions · Atelier collectif</text><text x="60" y="920" font-family="Arial,sans-serif" font-size="21">Accès prévu gratuit, sur inscription</text><text x="60" y="1020" font-family="Arial,sans-serif" font-size="17" font-weight="bold" fill="#92681d">ÉDITION PILOTE EN PRÉPARATION</text><text x="60" y="1055" font-family="Arial,sans-serif" font-size="15">Date et intervenants à confirmer · Inscriptions non ouvertes</text><text x="60" y="1120" font-family="Arial,sans-serif" font-size="13">Conception : Farès Benamar · Projet personnel de portfolio</text></g></svg>'''
(OUT/'affiche.svg').write_text(svg)

# A working concept page, with no fake registration form or invented confirmed date.
landing='''<!doctype html><html lang="fr"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Les Rencontres de L’Enjeu | Projet en préparation</title><style>*{box-sizing:border-box}body{margin:0;background:#f7f4ec;color:#202720;font:17px/1.8 Arial,sans-serif}main{max-width:1000px;padding:60px 28px;margin:auto}.brand{font:42px Georgia,serif}h1{font:clamp(42px,7vw,78px)/1.08 Georgia,serif;max-width:850px}h2{font:30px Georgia,serif;margin-top:50px}.status{color:#78531b;font-size:12px;font-weight:bold;letter-spacing:.1em}a{color:#202720;text-underline-offset:5px}:focus-visible{outline:3px solid #78531b;outline-offset:5px}.button{display:inline-block;border:1px solid #202720;border-radius:28px;padding:12px 20px;text-decoration:none;font-size:14px;margin-top:15px}.facts{display:flex;flex-wrap:wrap;gap:12px;margin:25px 0}.facts span{border:1px solid #cbd0c8;border-radius:25px;padding:7px 14px;font-size:13px}table{width:100%;border-collapse:collapse;font-size:15px}th,td{padding:14px 8px;text-align:left;border-bottom:1px solid #ccd0c7}footer{border-top:1px solid #ccd0c7;margin-top:55px;padding-top:22px;font-size:12px;color:#536056}</style></head><body><main><p class="brand">L’Enjeu.</p><p class="status">PROJET PERSONNEL · ÉDITION PILOTE EN PRÉPARATION</p><h1>Les Rencontres<br>de L’Enjeu.</h1><p>Comprendre l’information à l’heure des réseaux sociaux et de l’IA.</p><div class="facts"><span>Avril 2027 envisagé</span><span>Paris / Île-de-France</span><span>Accès prévu gratuit</span></div><p>Une table ronde, vos questions et un atelier collectif pour prendre du recul sur les informations qui circulent. La date exacte, le lieu et les intervenants restent à confirmer.</p><a class="button" href="mailto:contact.pro.faresb@gmail.com?subject=Les%20Rencontres%20de%20L%27Enjeu%20-%20informations">Recevoir les informations de préparation</a><h2>Au programme envisagé</h2><table><thead><tr><th>Horaire</th><th>Séquence</th></tr></thead><tbody>ROWS</tbody></table><h2>Pourquoi cette rencontre ?</h2><p>Identifier les sources, rechercher les preuves et remettre les informations en contexte. Un format ouvert à la discussion, pour les lecteurs du média et les personnes intéressées par les usages de l’information.</p><h2>Informations pratiques</h2><p>Première édition envisagée en avril 2027, de 18 h 30 à 21 h 30, à Paris ou en Île-de-France. La jauge de travail est de 80 participants. L’ouverture des inscriptions dépend de la confirmation du lieu et de l’organisation.</p><p>Les informations pratiques et l’ouverture des inscriptions seront annoncées après confirmation de l’organisation.</p><footer>Page de conception réalisée pour le portfolio de Farès Benamar. Aucun événement passé ni résultat de fréquentation n’est revendiqué.</footer></main></body></html>'''
landing=landing.replace('ROWS',''.join('<tr><td>'+a+' – '+b+'</td><td>'+escape(c)+'</td></tr>' for a,b,c,*_ in RUN if a>='18:30' and a<='21:05'))
(OUT/'08-page-presentation.html').write_text(landing)
(OUT/'LIRE-MOI.md').write_text('''# Les Rencontres de L’Enjeu

Projet personnel de conception, préparé par Farès Benamar le 13 octobre 2026. Première édition envisagée en avril 2027 ; date, lieu, équipe, intervenants, budget et financement non confirmés. Aucun événement réalisé ni résultat acquis n’est revendiqué.

## Contenu du dossier

- 01 : dossier de conception (PDF et DOCX modifiable).
- 02 : organisation, responsabilités, rétroplanning, conducteur et risques (PDF et DOCX).
- 03 : textes de communication, calendrier et indicateurs (PDF et DOCX).
- 04 : affiche A3 (PDF), avec aperçu vectoriel affiche.svg.
- 05 : invitation A5 (PDF).
- 06 : programme public envisagé (PDF).
- 07 : classeur de pilotage XLSX, dix feuilles avec hypothèses, formules et résultats vierges.
- 08 : page de présentation HTML autonome, sans formulaire ni collecte active.
- sources-projet.json : données et textes sources.

## Repères cohérents

80 participants visés, salle de 100 personnes au minimum, équipe de 8 personnes maximum. Budget prévisionnel TTC : 2 100 € de base + 210 € de réserve = 2 310 €. Aucun devis ni financement acquis. La référence du 8 avril 2027 dans le classeur sert uniquement au calendrier de travail ; les supports publics indiquent avril 2027 sans jour confirmé.

Les documents de communication comportant des champs entre crochets sont des modèles à compléter avant diffusion. Aucun destinataire n’a été contacté, aucune inscription ouverte et aucun prestataire engagé.

## Réutilisation

Dans le dépôt Portfolio, le script scripts/build-rencontres.py régénère les documents avec reportlab, python-docx et openpyxl. Les PDF sont prêts à consulter et les DOCX/XLSX permettent des ajustements. Les valeurs modifiées dans le XLSX doivent aussi être reportées dans les documents de conception avant toute diffusion.
''')
with zipfile.ZipFile(OUT/'dossier-complet.zip','w',zipfile.ZIP_DEFLATED) as z:
 for path in sorted(OUT.iterdir()):
  if path.is_file() and path.name!='dossier-complet.zip':z.write(path,path.name)
print('Generated',len(list(OUT.iterdir())),'files; budget:',TOTAL,'EUR; public period:',E['period'])
