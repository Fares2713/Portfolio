#!/usr/bin/env python3
"""Rebuild the fictional decision dossier from contenus.json and the JS decision model.
Requires Python reportlab and python-docx; Node runs the shared pure decision model.
"""
import csv, html, io, json, os, re, subprocess, zipfile
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, KeepTogether
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.colors import HexColor, white
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'assets/dossier-decisions';OUT.mkdir(parents=True,exist_ok=True)
DATA=json.loads((ROOT/'dossier-decisions/contenus.json').read_text())
js="""import fs from 'node:fs';import {decide,toMarkdown} from './dossier-decisions/moteur.mjs';const d=JSON.parse(fs.readFileSync('dossier-decisions/contenus.json'));console.log(JSON.stringify(d.cases.map(c=>({base:decide(d,{cas:c.id}),pending:decide(d,{cas:c.id,contrainte:'incomplet'}),access:decide(d,{cas:c.id,contrainte:'accessibilite'}),markdown:toMarkdown(d,decide(d,{cas:c.id}))}))));"""
MODELS=json.loads(subprocess.check_output([os.environ.get('PORTFOLIO_NODE','node'),'--input-type=module','-e',js],cwd=ROOT,text=True))
FONT=Path('/usr/share/fonts/truetype/dejavu')
pdfmetrics.registerFont(TTFont('Dossier',str(FONT/'DejaVuSans.ttf')))
pdfmetrics.registerFont(TTFont('DossierBold',str(FONT/'DejaVuSans-Bold.ttf')))
pdfmetrics.registerFont(TTFont('DossierSerif',str(FONT/'DejaVuSerif.ttf')))
BLUE=HexColor('#173ec4');INK=HexColor('#172c2a');MUTED=HexColor('#52615e');LINE=HexColor('#d6dcd4')
ST={
'body':ParagraphStyle('body',fontName='Dossier',fontSize=8.8,leading=13,textColor=INK,spaceAfter=7),
'small':ParagraphStyle('small',fontName='Dossier',fontSize=7.5,leading=11,textColor=MUTED,spaceAfter=8),
'h1':ParagraphStyle('h1',fontName='Dossier',fontSize=24,leading=29,textColor=INK,spaceAfter=15),
'h2':ParagraphStyle('h2',fontName='DossierBold',fontSize=11,leading=15,textColor=INK,spaceBefore=10,spaceAfter=6,keepWithNext=True),
'eyebrow':ParagraphStyle('eyebrow',fontName='DossierBold',fontSize=8,leading=12,textColor=BLUE,spaceAfter=14),
'quote':ParagraphStyle('quote',fontName='DossierSerif',fontSize=11,leading=16,textColor=BLUE,spaceBefore=8,spaceAfter=12)}
def P(text,style='body'):return Paragraph(html.escape(text).replace('\n','<br/>'),ST[style])
def H(label):return P(label,'h2')
def bullets(items):return [P('— '+x) for x in items]
def page(c,doc):
 c.saveState();w,h=doc.pagesize;c.setFillColor(BLUE);c.rect(0,h-7,w,7,fill=1,stroke=0)
 c.setStrokeColor(LINE);c.line(45,43,w-45,43);c.setFillColor(MUTED);c.setFont('Dossier',7)
 c.drawString(45,30,'FARÈS BENAMAR / EXERCICE FICTIF / V'+DATA['version']);c.drawRightString(w-45,30,str(doc.page));c.restoreState()
def pdf(name,story):
 doc=SimpleDocTemplate(str(OUT/name),pagesize=(595.28,841.89),leftMargin=45,rightMargin=45,topMargin=43,bottomMargin=60,title='Dossier de décisions — Farès Benamar',author=DATA['author'])
 doc.build(story,onFirstPage=page,onLaterPages=page)
def case_story(model,lead=True):
 r=model['base'];c=r['c'];p=model['pending'];a=model['access']
 story=[P('CAS '+c['number']+' / EXERCICE DE CONCEPTION','eyebrow'),P(c['title'],'h1'),P(c['tag'],'quote'),P(DATA['notice'],'small'),H('Le problème'),P(c['context']),P(c['problem']),H('Ce que je vérifie')]
 for f in r['verification']:
  story += [P(f['item'],'h2'),P(f['source']+' · '+f['owner']+' · '+f['state'],'small'),P(f['detail'])]
 story += [PageBreak(),P('CAS '+c['number']+' / LE PARCOURS DE RÉFÉRENCE','eyebrow'),P('Une réponse et un ordre de publication.','h1'),P('Public : usagers · Canal : e-mail · Contrainte : deux heures pour agir.','small'),H('Mon arbitrage'),P(r['title']),P(r['reason']),H('La condition de diffusion'),P(r['gate']),H('Le message central'),P(r['message'],'quote'),H('L’ordre de publication')]
 for t in r['timeline']:story += [P(t['when'],'h2'),P(t['text'])]
 story += [H('Ce que je refuse'),P(r['refusal']),PageBreak(),P('CAS '+c['number']+' / CONTENUS ET VARIANTES','eyebrow'),P('Ce qui change avec la contrainte.','h1'),H('Exemple d’e-mail — '+r['headline']),P(r['restriction'],'small'),P(r['body']),H('Si une information reste à confirmer'),P(p['title']),P(p['message'],'quote'),P(p['gate']),H('Si le public est éloigné du numérique'),P(a['access']),H('Évaluer, puis ajuster')]
 for m in r['metrics']:story += [P(m['name'],'h2'),P(m['method']),P(m['rule'],'small')]
 story += [P('Ces critères décrivent un protocole proposé, pas des résultats mesurés. Les variantes complètes sont accessibles dans le dossier interactif.','small')]
 return story
GUIDE=[P('LE DOSSIER DE DÉCISIONS / FARÈS BENAMAR','eyebrow'),Spacer(1,35),P('Avant de publier,\nil faut choisir.','h1'),P('Trois situations. Des faits à vérifier. Des décisions à expliquer.','quote'),P(DATA['notice']),Spacer(1,18),H('Comment lire ce dossier'),P('Les fiches développent un parcours de référence : usagers, e-mail, deux heures pour agir. Elles montrent aussi ce qui change lorsqu’une validation manque ou que le public est peu à l’aise avec le numérique.'),P('L’interface propose 81 parcours : trois situations, trois publics, trois canaux et trois contraintes. Certaines demandes de canal sont réorientées, avec une explication : un réseau social ouvert ne remplace pas un briefing interne.'),H('Ce que cet exercice rend visible'),P('Définir une prochaine étape utile, nommer les incertitudes, adapter le support au public et organiser une réponse cohérente entre communication et accueil.'),H('Ce qu’il ne prétend pas démontrer'),P('Il ne s’agit pas d’une mission réalisée pour une institution, ni d’une expérience de gestion de crise effective. Les résultats ne sont pas mesurés. Les exemples sont des maquettes de conception.'),P('Version '+DATA['version']+' · Sources éditoriales : contenus.json et moteur.mjs. Les PDF, le DOCX et la grille sont produits à partir de cette même base.','small')]
METHOD=[P('LA MÉTHODE / EXERCICE DE CONCEPTION','eyebrow'),P('Une décision doit pouvoir être expliquée.','h1'),P(DATA['notice'],'small'),H('1. Poser le contrat de communication'),P('Identifier le problème pratique, le public prioritaire, l’action attendue et le responsable de validation. Distinguer une information, une consultation et une rectification : elles n’appellent pas les mêmes preuves ni la même diffusion.'),H('2. Vérifier avant de formuler'),P('Créer une référence commune avec les faits établis, les éléments en attente et leur propriétaire. Une pièce fictive du cas représente un contrôle à réaliser, pas une validation réelle. Sans validation, différer la promesse concernée ; expliquer ce qui est connu si une réponse est nécessaire.'),H('3. Garder une chaîne de publication'),P('Aligner le responsable de service et l’accueil avant les relais publics. Le support choisi n’est pas nécessairement le premier acte de communication. Désigner qui met à jour la référence et qui retire les versions devenues obsolètes.'),H('4. Choisir un canal pour son rôle'),P('L’e-mail adresse un public identifié ; il nécessite une liste pertinente et un cadre d’envoi vérifié. Le réseau social oriente et répond, sans devenir le seul point de référence. L’affichage aide sur place, à condition de prévoir le retrait des versions périmées.'),H('5. Faire une place aux personnes hors du canal'),P('Donner une réponse orale et une voie d’accès sans smartphone. Un QR code ou un lien peut compléter le message ; il ne doit pas conditionner sa compréhension.'),H('6. Évaluer sans inventer de performance'),P('Les indicateurs servent une décision : une confusion conduit à réécrire ; une reprise divergente conduit à réaligner ; une correction invisible conduit à intervenir dans l’espace concerné. Les observations sont à réaliser, et les champs de résultat restent vides.'),PageBreak(),P('LES RÈGLES DE COHÉRENCE','eyebrow'),P('Une référence. Plusieurs supports.','h1')]
for a in DATA['audiences'].values():METHOD += [H(a['label']),P(a['need']),P(a['action'])]
METHOD += [H('Traçabilité et données personnelles'),P('Consigner le sujet, la pièce de référence, la validation et la version du message. Ne pas exporter d’adresses, de captures nominatives ni de témoignages identifiants dans le dossier de suivi. Le partage du parcours interactif ne contient que les quatre choix ; il ne collecte aucune donnée de recruteur.'),H('Réagir à une erreur de notre propre message'),P('Corriger la référence, rendre la correction explicite dans l’espace concerné, informer l’accueil et les relais, puis retirer les supports dépassés. Une correction ne doit pas être masquée par un simple remplacement silencieux.'),H('Rôle de l’auteur'),P('Farès Benamar : conception des cas, proposition de méthode, rédaction des maquettes, organisation de l’évaluation et réalisation du dossier interactif. L’exercice ne revendique aucune validation ou collaboration par l’institution fictive.')]
pdf('methode-editoriale.pdf',METHOD.copy())
allstory=GUIDE.copy()
for c,m in zip(DATA['cases'],MODELS):
 story=case_story(m)
 pdf(c['number']+'-'+c['id']+'.pdf',story.copy())
 allstory += [PageBreak()]+story
 (OUT/(c['number']+'-'+c['id']+'.md')).write_text(m['markdown'])
pdf('dossier-decisions.pdf',allstory)
# Editable source for reviewers: real paragraphs, not screenshots of the PDF.
doc=Document();doc.core_properties.author=DATA['author'];normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
doc.add_heading('Dossier de décisions',0);doc.add_paragraph('Farès Benamar · Exercice de conception · Version '+DATA['version']);doc.add_paragraph(DATA['notice'])
for model in MODELS:
 doc.add_page_break()
 for line in model['markdown'].split('\n'):
  if line.startswith('# '):doc.add_heading(line[2:],1)
  elif line.startswith('## '):doc.add_heading(line[3:],2)
  elif line.startswith('- '):doc.add_paragraph(line[2:],style='List Bullet')
  elif line.strip():doc.add_paragraph(line)
 doc.add_heading('Variante : validation manquante',2);doc.add_paragraph(model['pending']['title']);doc.add_paragraph(model['pending']['message']);doc.add_paragraph(model['pending']['gate'])
doc.save(OUT/'dossier-decisions.docx')
with (OUT/'grille-evaluation.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.writer(f,delimiter=';');writer.writerow(['Statut','Cas','Critère','Méthode proposée','Décision si difficulté','Observation réelle (à renseigner)','Date (à renseigner)','Responsable (à renseigner)'])
 for m in MODELS:
  for metric in m['base']['metrics']:writer.writerow(['EXERCICE FICTIF — AUCUN RÉSULTAT MESURÉ',m['base']['c']['title'],metric['name'],metric['method'],metric['rule'],'','',''])
guide='''# Dossier de décisions — guide de lecture\n\n'''+DATA['notice']+'''\n\nConception : Farès Benamar. Version '''+DATA['version']+'''.\n\n## Les fichiers\n\n- dossier-decisions.pdf : les trois cas détaillés et leur parcours de référence.\n- 01-service.pdf, 02-controverse.pdf, 03-rectification.pdf : fiches indépendantes.\n- methode-editoriale.pdf : principes communs et règles de cohérence.\n- dossier-decisions.docx : version modifiable des parcours de référence et des réponses en attente de validation.\n- grille-evaluation.csv : neuf critères, méthodes et décisions ; champs d’observation volontairement vides. Séparateur : point-virgule ; encodage UTF-8 avec BOM.\n- 01-service.md, 02-controverse.md, 03-rectification.md : parcours de référence en texte.\n- sources/contenus.json et sources/moteur.mjs : base éditoriale et règles de composition.\n\n## Explorer\n\nOuvrir dossier-decisions/ depuis le portfolio servi par HTTP. Choisir une situation, un public, un canal et une contrainte puis afficher l’arbitrage. Le lien du parcours conserve les quatre choix. Un export Markdown ou l’impression en PDF reprend la sélection affichée.\n\nLes 81 parcours résultent de règles éditoriales explicites, pas d’une génération par une IA lors de la visite. Les supports partagent les faits du cas. Certaines demandes sont réorientées, notamment le choix d’un canal public pour un briefing interne.\n\n## Référence éditoriale\n\nLes fiches décrivent le parcours usagers / e-mail / deux heures pour agir. Une validation manquante change le statut du message, parfois jusqu’à différer une annonce. Aucun responsable réel n’a validé les pièces fictives.\n'''
(OUT/'guide-lecture.md').write_text(guide)
with zipfile.ZipFile(OUT/'dossier-complet.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in sorted(OUT.iterdir()):
  if f.suffix!='.zip':z.write(f,f.name)
 z.write(ROOT/'dossier-decisions/contenus.json','sources/contenus.json');z.write(ROOT/'dossier-decisions/moteur.mjs','sources/moteur.mjs')
# Keep non-interactive summaries and case labels aligned with the source.
p=ROOT/'dossier-decisions/index.html';page=p.read_text()
blocks={'CASES':'\n'.join('<label class="case-option"><input type="radio" name="cas" form="choices" value="'+c['id']+'" '+('checked' if i==0 else '')+'><span class="case-number">'+c['number']+'</span><strong>'+html.escape(c['title'])+'</strong><small>'+html.escape(c['tag'])+'</small></label>' for i,c in enumerate(DATA['cases'])),
'FALLBACK':'\n'.join('<article><h3>'+html.escape(c['title'])+'</h3><p>'+html.escape(c['context'])+'</p><p><strong>Mon choix :</strong> '+html.escape(c['decision'])+'</p><p>'+html.escape(c['message'])+'</p><a href="../assets/dossier-decisions/'+c['number']+'-'+c['id']+'.pdf">Lire la fiche complète (PDF)</a></article>' for c in DATA['cases']),
'DOWNLOADS':'\n'.join('<article class="download-card"><p class="eyebrow">CAS '+c['number']+' / FICHE PDF</p><h3>'+html.escape(c['title'])+'</h3><p>'+html.escape(c['learning'])+'</p><a href="../assets/dossier-decisions/'+c['number']+'-'+c['id']+'.pdf">Consulter la fiche <svg class="icon" viewBox="0 0 24 24" width="15" height="15" fill="none" stroke="currentColor" stroke-width="1.5" stroke-linecap="round" stroke-linejoin="round" aria-hidden="true"><path d="M5 19 19 5M5 5h14v14"/></svg></a></article>' for c in DATA['cases'])}
for name,block in blocks.items():page=re.sub(r'<!-- '+name+r':START -->.*?<!-- '+name+r':END -->','<!-- '+name+':START -->\n'+block+'\n<!-- '+name+':END -->',page,flags=re.S)
p.write_text(page)
print('Generated dossier PDF, 3 case PDFs, method PDF, editable DOCX, evaluation CSV, 4 Markdown files and ZIP.')
