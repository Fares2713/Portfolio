#!/usr/bin/env python3
"""Build the event management workbook from the shared dossier assumptions."""
from pathlib import Path
from datetime import date,time
import json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule,ColorScaleRule,DataBarRule
from openpyxl.chart import BarChart,DoughnutChart,Reference
from openpyxl.chart.series import DataPoint
from openpyxl.workbook.properties import CalcProperties
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT=Path(__file__).resolve().parents[1]/'assets/rencontres-enjeu'
S=json.loads((OUT/'sources-projet.json').read_text())
INK='202A35';GOLD='B88B35';PAPER='F7F4EC';BLUE='173EC4';GREEN='18715A';LINE='DFE4E8';RED='B73E43'
EUR='#,##0.00 "€"';PCT='0.0%';DATE='dd/mm/yyyy'
wb=Workbook();wb.remove(wb.active);wb.calculation=CalcProperties(fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
wb.properties.title='Les Rencontres de L’Enjeu | Pilotage opérationnel';wb.properties.creator='Farès Benamar';wb.properties.description='Scénario de conception pour avril 2027. Aucun engagement ni résultat acquis.'
ends={}
def sheet(name,title,desc,headers,widths,rows):
 w=wb.create_sheet(name);n=len(headers)
 w.sheet_view.showGridLines=False;w.sheet_view.zoomScale=85
 w.row_dimensions[1].height=24;w.row_dimensions[2].height=24
 w.merge_cells(start_row=1,start_column=1,end_row=2,end_column=n);c=w.cell(1,1,title);c.font=Font(name='Aptos Display',size=24,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=INK);c.alignment=Alignment(vertical='center')
 for row in w.iter_rows(min_row=1,max_row=2,max_col=n):
  for c in row:c.fill=PatternFill('solid',fgColor=INK)
 w.merge_cells(start_row=4,start_column=1,end_row=5,end_column=n);c=w.cell(4,1,desc);c.font=Font(name='Aptos',size=10,color='596573');c.alignment=Alignment(wrap_text=True,vertical='center');w.row_dimensions[4].height=22;w.row_dimensions[5].height=18
 w.cell(6,1,'← TABLEAU DE BORD').hyperlink="#'Tableau de bord'!A1";w.cell(6,1).font=Font(size=9,color=BLUE,underline='single')
 for j,h in enumerate(headers,1):c=w.cell(7,j,h);c.font=Font(name='Aptos',size=10,bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor=GOLD);c.alignment=Alignment(wrap_text=True,vertical='center')
 w.row_dimensions[7].height=34
 for i,row in enumerate(rows,8):
  w.row_dimensions[i].height=43
  for j,v in enumerate(row,1):
   c=w.cell(i,j,v);c.font=Font(name='Aptos',size=10,color=GREEN if isinstance(v,str) and v.startswith('=') else INK);c.fill=PatternFill('solid',fgColor='F2F5F7' if i%2==0 else 'FFFFFF');c.alignment=Alignment(wrap_text=True,vertical='top');c.border=Border(bottom=Side(style='hair',color=LINE))
 for j,width in enumerate(widths,1):w.column_dimensions[L(j)].width=width
 end=7+len(rows);ends[name]=end
 w.freeze_panes='C8';w.auto_filter.ref=f'A7:{L(n)}{end}'
 if rows:
  t=Table(displayName='T'+str(len(wb.worksheets)),ref=w.auto_filter.ref);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False);w.add_table(t)
 w.sheet_properties.pageSetUpPr.fitToPage=True;w.page_setup.orientation='landscape';w.page_setup.paperSize=w.PAPERSIZE_A3;w.page_setup.fitToWidth=1;w.page_setup.fitToHeight=0;w.print_title_rows='1:7';w.print_options.horizontalCentered=True;w.print_area=f'A1:{L(n)}{end}'
 w.oddHeader.center.text='Les Rencontres de L’Enjeu | Conception · Avril 2027';w.oddFooter.left.text='Date, lieu et intervenants à confirmer';w.oddFooter.right.text='Page &P / &N'
 return w

def inputs(w,cols,start=8,end=None):
 for row in range(start,(end or ends[w.title])+1):
  for col in cols:
   c=w.cell(row,col)
   if c.data_type!='f':c.font=Font(name='Aptos',size=10,color=BLUE);c.fill=PatternFill('solid',fgColor='EAF0FF')

def fmt(w,cols,pattern,start=8,end=None):
 for row in range(start,(end or ends[w.title])+1):
  for col in cols:w.cell(row,col).number_format=pattern

def valid(w,col,formula,start=8,end=None,kind='list',op=None,formula2=None):
 v=DataValidation(type=kind,formula1=formula,formula2=formula2,operator=op,allow_blank=True);v.errorTitle='Valeur à vérifier';v.error='Utilisez une valeur autorisée.';v.showErrorMessage=True;v.errorStyle='stop';w.add_data_validation(v);v.add(f'{col}{start}:{col}{end or ends[w.title]}')

def statuses(w,col):
 valid(w,col,'Statuts');w.conditional_formatting.add(f'{col}8:{col}{ends[w.title]}',FormulaRule(formula=[f'{col}8="Terminé"'],fill=PatternFill('solid',fgColor='DBF0E7'),font=Font(color=GREEN)))
 w.conditional_formatting.add(f'{col}8:{col}{ends[w.title]}',FormulaRule(formula=[f'{col}8="Bloqué"'],fill=PatternFill('solid',fgColor='FCE5E5'),font=Font(color=RED)))

sheet('Tableau de bord','LES RENCONTRES / PILOTAGE','Édition pilote envisagée en avril 2027 · Hypothèses de travail, sans événement réalisé.', ['Repère','Valeur','Lecture'],[37,28,65],[])
params=sheet('Paramètres','01 / HYPOTHÈSES COMMUNES','Modifier les cellules bleues. Les dates, budgets et jauges des autres onglets se recalculent depuis cette feuille.', ['Paramètre','Valeur','Règle / justification'],[38,27,91],[
 ['Date de référence',date(2027,4,8),'Référence de calcul uniquement ; le jour exact n’est pas confirmé.'],
 ['Date de revue',date(2026,10,13),'Date saisie, pour rendre les alertes reproductibles. À actualiser à chaque revue.'],
 ['Participants visés',80,'Objectif de conception ; ne constitue pas une fréquentation acquise.'],
 ['Capacité minimale du lieu',100,'À vérifier avec le gestionnaire du lieu.'],
 ['Effectif équipe maximal',8,'Enveloppe de travail, intervenants inclus.'],
 ['Réserve budgétaire',.1,'Appliquée uniquement aux lignes de budget prévisionnel.'],
 ['Plafond de travail TTC',2500,'Limite de scénario à valider avant engagement.'],
 ['Nombre de questionnaires envoyés',None,'Résultat à renseigner après envoi effectif.'],
 ['Nombre de réponses reçues',None,'Résultat à renseigner après réception.'],
 ['Inscriptions confirmées',None,'Total agrégé ; aucune liste nominative dans ce classeur.'],
 ['Entrées uniques constatées',None,'Comptage sans double entrée, après réalisation.'],
 ['Heure accueil public',time(18,30),'Horaire envisagé ; conducteur de référence.'],
 ['Heure fin du public',time(21,30),'Horaire envisagé.'],
 ['Coût cible par participant',"='Budget'!H36/B10",'Coût prévisionnel, calculé sur la jauge cible.'],
 ])
inputs(params,[2]);fmt(params,[2],DATE,8,9);params['B13'].number_format=PCT;params['B14'].number_format=EUR;fmt(params,[2],'hh:mm',19,20);params['B21'].number_format=EUR
valid(params,'B','1',8,9,kind='date',op='greaterThan');valid(params,'B','1',10,12,kind='whole',op='greaterThanOrEqual');valid(params,'B','0',13,13,kind='decimal',op='between',formula2='1');valid(params,'B','0',14,14,kind='decimal',op='greaterThanOrEqual');valid(params,'B','0',15,18,kind='whole',op='greaterThanOrEqual')

# Detailed cost lines keep the eight reference envelopes and the 2 100 EUR base.
costs=[
 ('Lieu','Location de salle',1,500,'Forfait salle ; à comparer par devis'),('Lieu','Nettoyage / restitution',1,60,'Hypothèse si non inclus'),('Lieu','Mobilier complémentaire',1,40,'Option selon équipement du lieu'),
 ('Technique','Technicien / régie',1,250,'Installation, essais et conduite'),('Technique','Deux micros et sonorisation',1,150,'Forfait de location'),('Technique','Câblage / adaptateurs',1,50,'Lot de dépannage'),
 ('Intervenants','Intervention 1',1,100,'Enveloppe proposée ; conditions non convenues'),('Intervenants','Intervention 2',1,100,'Enveloppe proposée ; conditions non convenues'),('Intervenants','Modération',1,100,'Enveloppe proposée ; conditions non convenues'),
 ('Collation','Boissons',80,1,'Par participant cible'),('Collation','Collation',80,1.5,'Par participant cible'),('Collation','Consommables',80,.5,'Par participant cible'),
 ('Supports','Affiches A3',10,2,'Hypothèse impression locale'),('Supports','Invitations A5',50,.4,'Hypothèse tirage limité'),('Supports','Programmes / repères',80,.5,'Par participant cible'),('Supports','Signalétique accueil',1,10,'Lot'),
 ('Administration','Assurance événement',1,80,'À vérifier avec la structure porteuse'),('Administration','Frais administratifs',1,40,'Provision de scénario'),
 ('Accessibilité','Adaptations / accompagnement',1,120,'Provision ; besoins à vérifier avec le lieu'),
 ('Documentation','Photographie / documentation',1,180,'Option à arbitrer avant engagement'),
 ]
rows=[]
for i,(cat,label,q,pu,note) in enumerate(costs,8):
 quantity="='Paramètres'!$B$10" if label in ['Boissons','Collation','Consommables','Programmes / repères'] else q
 rows.append([f'B{i-7:02}',cat,label,'participant' if isinstance(quantity,str) else 'unité',quantity,pu,note,f'=ROUND(E{i}*F{i},2)',None,None,f'=IF(COUNT(I{i})=0,"",I{i}-H{i})','À chiffrer'])
b=sheet('Budget','02 / BUDGET PRÉVISIONNEL TTC','20 lignes de coût reliées à la jauge. Devis et validations à renseigner ; aucun montant n’est engagé.', ['ID','Enveloppe','Poste','Unité','Quantité','PU TTC','Hypothèse','Prévision TTC','Devis TTC','Référence devis','Écart devis','Validation'],[10,20,34,16,13,17,52,19,19,25,19,20],rows)
inputs(b,[5,6,9,10,12]);fmt(b,[6,8,9,11],EUR);valid(b,'E','0',kind='decimal',op='greaterThanOrEqual');valid(b,'F','0',kind='decimal',op='greaterThanOrEqual');valid(b,'L','"À chiffrer,Devis reçu,À arbitrer,Validé"')
for r in range(33,40):b.row_dimensions[r].height=26
for r,label,value in [(33,'BASE PRÉVISIONNELLE','=SUM(H8:H27)'),(34,'TAUX DE RÉSERVE',"='Paramètres'!B13"),(35,'RÉSERVE TTC','=ROUND(H33*H34,2)'),(36,'TOTAL PRÉVISIONNEL TTC','=H33+H35'),(37,'PLAFOND TTC',"='Paramètres'!B14"),(38,'MARGE SUR PLAFOND','=H37-H36'),(39,'COÛT PAR PARTICIPANT',"=IFERROR(H36/'Paramètres'!B10,0)")]:
 b.cell(r,7,label);c=b.cell(r,8,value);c.font=Font(size=12,bold=True,color=GREEN);c.number_format=PCT if r==34 else EUR
b.print_area='A1:L39';b.conditional_formatting.add('H38',FormulaRule(formula=['H38<0'],fill=PatternFill('solid',fgColor='FCE5E5')))
fin=sheet('Financements','03 / PLAN DE FINANCEMENT','Les intentions ne sont pas des ressources acquises. Saisir les accords et encaissements uniquement sur preuve.', ['ID','Source envisagée','Nature','Hypothèse TTC','Accord TTC','Encaissé TTC','Statut','Preuve / référence','Responsable','Prochaine action'],[10,32,23,19,19,19,23,35,23,50],[
 ['F01','Structure porteuse','Contribution interne',1500,None,None,'À rechercher',None,'Responsable projet','Valider la capacité de financement.'],
 ['F02','Soutien à identifier','Contribution complémentaire',"=MAX(0,'Budget'!H36-D8)",None,None,'À rechercher',None,'Responsable projet','Identifier une ressource ou réduire les options.'],
 ['F03','Apport en nature éventuel','Valorisation hors trésorerie',None,None,None,'À rechercher',None,'Logistique','Ne pas additionner au financement monétaire.']])
inputs(fin,[2,3,4,5,6,7,8,9,10]);fmt(fin,[4,5,6],EUR);valid(fin,'G','"À rechercher,Contact à préparer,En discussion,Accord écrit,Refusé"')
for col in 'DEF':valid(fin,col,'0',kind='decimal',op='greaterThanOrEqual')
for c,formula in [('D14','=SUM(D8:D9)'),('E14','=SUM(E8:E9)'),('F14','=SUM(F8:F9)'),('E15',"=MAX(0,'Budget'!H36-E14)"),('F15',"=MAX(0,'Budget'!H36-F14)")]:fin[c]=formula;fin[c].number_format=EUR
fin['A14']='RESSOURCES MONÉTAIRES';fin['A15']='RESTE À COUVRIR';fin.print_area='A1:J15'
dep=sheet('Dépenses','04 / ENGAGEMENTS ET PAIEMENTS','Registre vierge. Une ligne par pièce ; renseigner 0 si le coût réel est nul. Ne pas saisir ici les hypothèses du budget.', ['ID budget','Dépense','Fournisseur','Pièce / devis','Engagé TTC','Date engagement','Payé TTC','Date paiement','Statut','Pièce facture','Contrôle'],[13,32,27,25,18,19,18,19,24,25,52],[[None]*10+[f'=IF(COUNTA(A{r}:J{r})=0,"",IF(A{r}="","ID budget manquant",IF(COUNTIF(BudgetIDs,A{r})=0,"ID budget inconnu",IF(AND(ISNUMBER(G{r}),NOT(ISNUMBER(E{r}))),"Engagement à renseigner",IF(G{r}>E{r},"Paiement supérieur à engagement","À rapprocher")))))'] for r in range(8,48)])
inputs(dep,list(range(1,11)));fmt(dep,[5,7],EUR);fmt(dep,[6,8],DATE);valid(dep,'A','BudgetIDs');valid(dep,'I','"À valider,Engagé,Facturé,Payé,Annulé"');valid(dep,'E','0',kind='decimal',op='greaterThanOrEqual');valid(dep,'G','0',kind='decimal',op='greaterThanOrEqual')

# Expanded actionable baseline. Dependencies reference a single predecessor by ID.
tasks=[
 ('Cadrage','Valider objectifs et publics','Responsable projet',-56,-53,'','Dossier de cadrage','À faire',0),('Cadrage','Valider portage et plafond financier','Responsable projet',-53,-49,'T01','Décision écrite','À faire',0),
 ('Lieu','Définir critères de choix du lieu','Logistique',-49,-47,'T02','Grille de critères','À faire',0),('Lieu','Recueillir deux propositions de lieux','Logistique',-47,-43,'T03','Devis comparables','À faire',0),('Lieu','Vérifier accès, jauge et conditions','Logistique',-43,-42,'T04','Grille lieu complétée','À faire',0),('Lieu','Confirmer le lieu par écrit','Responsable projet',-42,-41,'T05','Accord du gestionnaire','À faire',0),
 ('Programme','Établir la liste des profils intervenants','Éditorial',-49,-44,'T02','Critères et profils','À faire',0),('Programme','Préparer les invitations intervenants','Éditorial',-44,-40,'T07','Invitations prêtes','À faire',0),('Programme','Obtenir accords et conditions','Éditorial',-40,-35,'T08','Accords écrits','À faire',0),('Programme','Valider fil conducteur et angle éditorial','Éditorial',-35,-31,'T09','Programme validé','À faire',0),
 ('Technique','Comparer les besoins et le devis technique','Logistique / régie',-41,-35,'T06','Fiche technique','À faire',0),('Technique','Valider assurance et adaptations','Logistique',-35,-29,'T11','Conditions confirmées','À faire',0),
 ('Communication','Affiche et invitation : conception graphique','Communication',-177,-177,'','Visuels A3 / A5 créés le 13 octobre 2026','Terminé',1),('Communication','Adapter visuels aux informations confirmées','Communication',-31,-29,'T10','Supports validés pour diffusion','À faire',0),('Communication','Préparer page et formulaire inscription','Communication / accueil',-35,-29,'T06','Parcours testé','À faire',0),('Communication','Valider ouverture des inscriptions','Responsable projet',-29,-28,'T15','Décision de diffusion','À faire',0),('Communication','Diffuser annonce et invitation','Communication',-28,-25,'T16','Publication et e-mail','À faire',0),('Communication','Contacter les relais pertinents','Communication',-25,-21,'T17','Suivi de diffusion','À faire',0),('Communication','Diffuser programme confirmé','Communication',-14,-12,'T18','Programme public','À faire',0),
 ('Conduite','Préparer conducteur et responsabilités','Responsable projet',-21,-14,'T10','Conducteur partagé','À faire',0),('Conduite','Briefer équipe et intervenants','Responsable projet',-14,-7,'T20','Brief et confirmations','À faire',0),('Accueil','Vérifier jauge et besoins d’accueil','Accueil',-7,-3,'T21','Organisation accueil','À faire',0),('Accueil','Envoyer rappel et accès','Communication',-2,-2,'T22','E-mail pratique','À faire',0),('Technique','Tester matériel et plan de secours','Régie',-2,-1,'T12','Tests validés','À faire',0),('Événement','Installer et contrôler le lieu','Logistique',0,0,'T24','Check-list ouverture','À faire',0),('Événement','Accueillir et conduire la soirée','Équipe',0,0,'T25','Journal de conduite','À faire',0),('Événement','Restituer salle et inventorier matériel','Logistique',0,0,'T26','Restitution vérifiée','À faire',0),('Bilan','Envoyer ressources et questionnaire','Communication',1,3,'T27','Message et questionnaire','À faire',0),('Bilan','Rapprocher dépenses et paiements','Responsable projet',3,5,'T28','Registre rapproché','À faire',0),('Bilan','Consolider fréquentation et retours','Responsable projet',5,7,'T29','Bilan renseigné','À faire',0),('Bilan','Décider les suites du projet','Responsable projet',7,7,'T30','Décision finale','À faire',0)]
rows=[]
for i,(phase,action,pilote,a,z,pre,liv,stat,prog) in enumerate(tasks,8):
 rows.append([f'T{i-7:02}',phase,action,pilote,a,z,pre,f"='Paramètres'!$B$8+E{i}",f"='Paramètres'!$B$8+F{i}",f'=I{i}-H{i}+1',stat,prog,liv,f'=IF(K{i}="Terminé","Terminé",IF(K{i}="Bloqué","Bloqué",IF(I{i}<\'Paramètres\'!$B$9,"En retard","Planifié")))',f'=IF(G{i}="","Sans dépendance",IFERROR(IF(H{i}>=INDEX($I$8:$I$38,MATCH(G{i},$A$8:$A$38,0)),"Cohérent","Chevauchement à arbitrer"),"ID inconnu"))'])
retro=sheet('Rétroplanning','05 / PLAN D’ACTION','31 tâches · dates relatives à la référence non confirmée · vérifier aussi les prérequis transverses dans Décisions.', ['ID','Phase','Action','Pilote proposé','Début J±','Fin J±','Prédécesseur','Début calculé','Fin calculée','Jours','Statut','Avancement','Livrable','Alerte échéance','Contrôle lien'],[10,21,57,26,13,13,15,18,18,12,21,17,42,23,32],rows)
retro['H20']=date(2026,10,13);retro['I20']=date(2026,10,13)
retro['E20']="=H20-'Paramètres'!$B$8";retro['F20']="=I20-'Paramètres'!$B$8"
inputs(retro,[4,5,6,7,11,12]);fmt(retro,[8,9],DATE);fmt(retro,[12],PCT);statuses(retro,'K');valid(retro,'L','0',kind='decimal',op='between',formula2='1');valid(retro,'G','TaskIDs');retro.conditional_formatting.add('L8:L38',DataBarRule(start_type='num',start_value=0,end_type='num',end_value=1,color=GREEN));retro.conditional_formatting.add('N8:N38',FormulaRule(formula=['N8="En retard"'],fill=PatternFill('solid',fgColor='FCE5E5')))

# Weekly Gantt, linked to the task register with conditional formatting.
g=sheet('Gantt','06 / CALENDRIER VISUEL','Une colonne par semaine. Doré = planifié, vert = terminé, rouge = blocage. Dates et statuts proviennent du rétroplanning.', ['ID','Action','Pilote','Début','Fin','Statut']+[f'Semaine {i+1}' for i in range(10)],[10,55,25,17,17,19]+[13]*10,[[f"='Rétroplanning'!{c}{r}" for c in ['A','C','D','H','I','K']]+[None]*10 for r in range(8,39)])
for j in range(7,17):g.cell(7,j,f"='Paramètres'!$B$8-56+{(j-7)*7}");g.cell(7,j).number_format='dd mmm';g.cell(7,j).font=Font(bold=True,color='FFFFFF');g.cell(7,j).fill=PatternFill('solid',fgColor=GOLD)
# Numeric dates in the Gantt header are not Excel table labels.
g._tables.clear();g.auto_filter.ref='A7:F38';fmt(g,[4,5],DATE)
for status,color in [('Terminé','18715A'),('Bloqué','B73E43')]:g.conditional_formatting.add('G8:P38',FormulaRule(formula=[f'AND(G$7<=$E8,G$7+6>=$D8,$F8="{status}")'],fill=PatternFill('solid',fgColor=color),stopIfTrue=True))
g.conditional_formatting.add('G8:P38',FormulaRule(formula=['AND(G$7<=$E8,G$7+6>=$D8)'],fill=PatternFill('solid',fgColor=GOLD)))

run=sheet('Déroulé','07 / CONDUCTEUR DU JOUR J','Horaire de référence envisagé · aucun intervenant confirmé · durée et enchaînement contrôlés par formule.', ['N°','Début','Fin','Durée min','Séquence','Pilote proposé','Préparation','Vigilance / repli','Contrôle horaire'],[9,13,13,15,33,30,50,67,27],[[i,time.fromisoformat(a),time.fromisoformat(z),f'=ROUND((C{r}-B{r})*1440,0)',seq,owner,prep,risk,('Début' if i==1 else f'=IF(B{r}=C{r-1},"Continu",IF(B{r}<C{r-1},"Chevauchement","Intervalle"))')] for i,(r,(a,z,seq,owner,prep,risk)) in enumerate(zip(range(8,22),S['run']),1)])
fmt(run,[2,3],'hh:mm');inputs(run,[2,3,6,7,8])
raci=sheet('RACI','08 / RÉPARTITION DES RESPONSABILITÉS','A = arbitrage (un seul par ligne), R = réalise, C = consulté, I = informé. Rôles proposés à attribuer à l’équipe.', ['Action','Responsable projet','Éditorial','Communication','Logistique','Accueil','Contrôle A'],[49,23,20,23,21,20,32],[list(row)+[f'=IF(COUNTIF(B{r}:F{r},"*A*")=1,"Un arbitrage identifié","Revoir responsabilité A")'] for r,row in enumerate(S['raci'],8)])
inputs(raci,[2,3,4,5,6]);
for c in 'BCDEF':valid(raci,c,'"A,R,C,I,A/R"')
com=sheet('Communication','09 / PLAN DE COMMUNICATION','Diffusion conditionnée à la validation du lieu, de la date, des intervenants et du parcours d’inscription.', ['ID','Échéance J±','Date calculée','Canal','Support','Public','Objectif','Pilote','Statut','URL / preuve de diffusion'],[10,16,18,23,42,29,48,23,22,47],[
 [f'C{i}',off,f"='Paramètres'!$B$8+B{r}",canal,support,public,obj,'Communication','À faire',None] for i,(r,(off,canal,support,public,obj)) in enumerate(zip(range(8,18),[
 (-28,'Page web','Annonce et inscription','Public cible','Donner des informations confirmées'),(-28,'E-mail','Invitation A5 et message','Lecteurs intéressés','Inviter après ouverture'),(-21,'Relais','Kit de partage','Associations et étudiants','Faire connaître le rendez-vous'),(-21,'Réseaux sociaux','Affiche A3 adaptée','Abonnés','Présenter le sujet'),(-14,'Réseaux / page','Programme confirmé','Public cible','Valoriser le format'),(-7,'Page web','Accès et informations pratiques','Inscrits','Faciliter la venue'),(-2,'E-mail','Rappel et désistement','Inscrits','Confirmer / libérer les places'),(1,'E-mail','Remerciement et ressources','Participants','Prolonger les échanges'),(3,'E-mail','Questionnaire facultatif','Participants contactés','Recueillir les retours'),(7,'Interne','Bilan factuel','Équipe','Décider les suites')]),1)])
inputs(com,[2,4,5,6,7,8,9,10]);fmt(com,[3],DATE);statuses(com,'I')
partners=sheet('Partenaires','10 / RELAIS ET PARTENAIRES','Pistes par type de structure ; aucun nom, contact ni accord inventé. Une proposition concrète doit être documentée.', ['ID','Type de structure','Intérêt recherché','Structure identifiée','Contact professionnel','Contribution envisagée','Statut','Prochaine action','Échéance','Preuve accord'],[10,30,43,33,34,45,23,48,19,38],[
 [f'P{i}',typ,interest,None,None,contrib,'À identifier','Qualifier une structure et préparer la prise de contact',None,None] for i,(typ,interest,contrib) in enumerate([
 ('Association étudiante','Public intéressé par les médias','Relais invitation'),('Bibliothèque / médiathèque','Lieu accessible et médiation','Lieu ou relais à discuter'),('Association éducation médias','Expertise sur les sources','Participation éditoriale'),('Média local','Ancrage et diffusion','Relais de communication'),('Structure culturelle locale','Ouverture à de nouveaux publics','Relais ou accueil'),('Soutien financier à identifier','Boucler le scénario financier','Contribution sous accord')],1)])
inputs(partners,list(range(4,11)));fmt(partners,[9],DATE);valid(partners,'G','"À identifier,À contacter,Contacté,En discussion,Accord écrit,Refusé"')
inter=sheet('Intervenants','11 / PROGRAMME ET INTERVENANTS','Profils recherchés, sans personne confirmée. Accord écrit et conditions indispensables avant annonce publique.', ['Rôle','Profil recherché','Nom confirmé','Structure','Statut','Conditions / rémunération','Accord écrit','Biographie validée','Disponibilité confirmée','Brief réalisé'],[23,60,30,28,25,48,23,25,28,22],[
 ['Intervenant 1','Journalisme / vérification des sources',None,None,'À identifier','Hypothèse de 100 € TTC à discuter','Non','Non','Non','Non'],['Intervenant 2','Recherche / usages de l’IA et des réseaux',None,None,'À identifier','Hypothèse de 100 € TTC à discuter','Non','Non','Non','Non'],['Modération','Animation de débat et médiation',None,None,'À identifier','Hypothèse de 100 € TTC à discuter','Non','Non','Non','Non']])
inputs(inter,list(range(3,11)));valid(inter,'E','"À identifier,Invité,En discussion,Confirmé,Indisponible"');
for c in 'GHIJ':valid(inter,c,'"Non,Oui,Sans objet"')
log=sheet('Logistique','12 / BESOINS ET INVENTAIRE','Quantités de travail ; la disponibilité, les réservations et l’attribution au lieu restent à vérifier.', ['ID','Zone','Besoin','Quantité prévue','Unité','Pilote','Mode envisagé','Statut','Point de contrôle'],[10,22,38,18,15,26,35,23,64],[
 [f'L{i}',zone,item,q,unit,owner,mode,'À faire',check] for i,(zone,item,q,unit,owner,mode,check) in enumerate([
 ('Salle','Places assises',"='Paramètres'!B10",'place','Logistique','Lieu','Garder circulations et places accessibles'),('Salle','Capacité totale minimale',"='Paramètres'!B11",'personne','Logistique','Gestionnaire','Inclure équipe dans le calcul des occupants'),('Audio','Micros intervenants',2,'micro','Régie','Location','Tester avant ouverture'),('Audio','Micro questions public',1,'micro','Régie','À vérifier avec le forfait','Éviter une dépense additionnelle non arbitrée'),('Audio','Enceintes',2,'enceinte','Régie','Lieu / location','Couverture de la salle et niveaux'),('Accueil','Tables accueil',2,'table','Accueil','Lieu','Flux, visibilité et accès'),('Accueil','Programmes',"='Paramètres'!B10",'exemplaire','Communication','Impression','Limiter le tirage au besoin'),('Atelier','Supports d’exercice',"='Paramètres'!B10",'support','Éditorial','Numérique / programme','Cas documenté et sources'),('Collation','Portions collation',"='Paramètres'!B10",'portion','Logistique','Approvisionnement','Contraintes alimentaires et déchets'),('Équipe','Personnel maximum',"='Paramètres'!B12",'personne','Responsable projet','À affecter','Intervenants inclus dans l’enveloppe'),('Sécurité','Repérage issues',1,'contrôle','Logistique','Gestionnaire','Consignes et référent du lieu'),('Technique','Jeu de câbles de secours',1,'lot','Régie','Location / équipe','Adaptateurs testés')],1)])
inputs(log,[4,6,7,8]);statuses(log,'H')
lieux=sheet('Lieux','13 / COMPARAISON DES LIEUX','Deux options à rechercher. Ne pas réserver avant vérification de la capacité, de l’accès et du coût global.', ['Option','Lieu identifié','Capacité vérifiée','Accessibilité vérifiée','Coût TTC','Technique incluse','Disponibilité','Conditions annulation','Accord écrit','Décision'],[15,35,21,26,20,38,24,47,23,24],[['Option A',None,None,'À vérifier',None,None,'À vérifier',None,'Non','À comparer'],['Option B',None,None,'À vérifier',None,None,'À vérifier',None,'Non','À comparer']])
inputs(lieux,list(range(2,11)));fmt(lieux,[5],EUR);valid(lieux,'D','"À vérifier,Oui,Non"');valid(lieux,'I','"Non,Oui"')
risks=sheet('Risques','14 / REGISTRE DES RISQUES','Échelle de 1 à 3. Score = probabilité × impact ; ≥ 6 = prioritaire. Évaluation initiale de conception.', ['ID','Risque','Probabilité','Impact','Score','Niveau','Prévention / repli','Pilote','Statut','Dernière revue'],[10,45,17,14,13,22,85,26,23,20],[
 [f'R{i}',risk,{'Fort':3,'Moyen':2,'Faible':1}[prob],{'Fort':3,'Moyen':2,'Faible':1}[imp],f'=C{r}*D{r}',f'=IF(E{r}>=6,"Prioritaire",IF(E{r}>=3,"À surveiller","Modéré"))',action,owner,'À faire',None] for i,(r,(risk,prob,imp,action,owner,_)) in enumerate(zip(range(8,16),S['risks']),1)])
inputs(risks,[3,4,7,8,9,10]);fmt(risks,[10],DATE);statuses(risks,'I');
for c in 'CD':valid(risks,c,'1',kind='whole',op='between',formula2='3')
risks.conditional_formatting.add('E8:E15',ColorScaleRule(start_type='num',start_value=1,start_color='DBF0E7',mid_type='num',mid_value=4,mid_color='FFF1CF',end_type='num',end_value=9,end_color='F3A7AB'))
dec=sheet('Décisions','15 / VALIDATIONS ET PRÉREQUIS','La diffusion et les engagements nécessitent des validations écrites ; aucune validation externe n’est acquise.', ['ID','Décision attendue','Pilote','Échéance J±','Date calculée','Statut','Preuve / décision','Conséquence si non validé'],[10,49,27,17,19,25,55,64],[
 [f'D{i}',topic,owner,off,f"='Paramètres'!$B$8+D{r}",'Non confirmé',None,impact] for i,(r,(topic,owner,off,impact)) in enumerate(zip(range(8,16),[
 ('Portage juridique et budget','Responsable projet',-49,'Aucun engagement financier'),('Financement disponible','Responsable projet',-42,'Arbitrer le périmètre ou reporter'),('Date et lieu','Responsable projet / logistique',-41,'Aucune annonce ferme'),('Intervenants et conditions','Éditorial',-35,'Ne pas annoncer de noms'),('Assurance et besoins techniques','Logistique',-29,'Ne pas confirmer ouverture au public'),('Supports et programme publics','Communication',-28,'Suspendre la diffusion'),('Inscription et gestion des données','Responsable projet / accueil',-28,'Ne pas ouvrir le formulaire'),('Ouverture le jour J','Responsable projet',0,'Suivre les consignes du gestionnaire')]),1)])
inputs(dec,[3,4,6,7,8]);fmt(dec,[5],DATE);valid(dec,'F','"Non confirmé,En cours,Validé,Refusé"')
control=sheet('Check-list','16 / CONTRÔLES OPÉRATIONNELS','Points à vérifier avec les responsables avant ouverture. Les preuves et contrôles restent à renseigner.', ['ID','Moment','Contrôle','Pilote','Statut','Preuve / commentaire'],[10,24,76,28,23,66],[
 [f'Q{i}',moment,check,owner,'À faire',None] for i,(moment,check,owner) in enumerate([
 ('Avant diffusion','Date, lieu, intervenants et modalités validés','Responsable projet'),('Avant diffusion','Texte sans information non confirmée','Communication'),('Avant diffusion','Formulaire testé et jauge paramétrée','Accueil'),('J-7','Liste et responsabilités équipe confirmées','Responsable projet'),('J-7','Accès et adaptations vérifiés avec le lieu','Logistique'),('J-2','Rappel envoyé et désistement possible','Communication'),('Avant ouverture','Capacité et circulations vérifiées','Logistique'),('Avant ouverture','Micros et secours testés','Régie'),('Avant ouverture','Issues et contacts du gestionnaire repérés','Logistique'),('Avant ouverture','Accueil et supports prêts','Accueil'),('Pendant','Entrées comptées sans doublons','Accueil'),('Pendant','Timing et incidents documentés','Responsable projet'),('Après','Matériel inventorié et salle restituée','Logistique'),('Après','Dépenses rapprochées des justificatifs','Responsable projet'),('Après','Retours consolidés sans résultats inventés','Responsable projet')],1)])
inputs(control,[4,5,6]);statuses(control,'E')
kpi=sheet('Indicateurs','17 / MESURE ET BILAN','Les résultats restent vides avant saisie. Les dépenses payées et engagements proviennent du registre Dépenses.', ['Indicateur','Valeur','Définition / lecture'],[44,27,85],[
 ['Participants visés',"='Paramètres'!B10",'Objectif de travail'],['Inscriptions confirmées','=IF(COUNT(\'Paramètres\'!B17)=0,"",\'Paramètres\'!B17)','Résultat à renseigner'],['Entrées uniques','=IF(COUNT(\'Paramètres\'!B18)=0,"",\'Paramètres\'!B18)','Résultat à renseigner'],['Taux de présence','=IF(OR(COUNT(B9:B10)<2,B9=0),"",B10/B9)','Entrées / inscriptions ; vide si pas de données'],['Questionnaires envoyés','=IF(COUNT(\'Paramètres\'!B15)=0,"",\'Paramètres\'!B15)','Résultat après envoi'],['Réponses reçues','=IF(COUNT(\'Paramètres\'!B16)=0,"",\'Paramètres\'!B16)','Résultat après réception'],['Taux de réponse','=IF(OR(COUNT(B12:B13)<2,B12=0),"",B13/B12)','Réponses / questionnaires envoyés'],['Budget prévisionnel TTC',"='Budget'!H36",'Hypothèses incluant la réserve'],['Engagements saisis','=IF(COUNT(\'Dépenses\'!E8:E47)=0,"",SUM(\'Dépenses\'!E8:E47))','Montants engagés saisis, pas le budget de conception'],['Paiements saisis','=IF(COUNT(\'Dépenses\'!G8:G47)=0,"",SUM(\'Dépenses\'!G8:G47))','Décaissements documentés ; pas un coût final avant clôture'],['Écart paiements - budget','=IF(COUNT(B17)=0,"",B17-B15)','Comparer avec prudence tant que des factures restent à payer'],['Taux tâches terminées','=COUNTIF(\'Rétroplanning\'!K8:K38,"Terminé")/COUNTA(\'Rétroplanning\'!A8:A38)','Avancement administratif du plan d’action'],['Réserve budgétaire',"='Budget'!H35",'Provision de conception'],['Solde engagements non payés','=IF(COUNT(B16:B17)<2,"",B16-B17)','Registre à rapprocher, non exhaustif avant clôture'],['Financement confirmé',"='Financements'!E14",'Somme des accords monétaires saisis ; aucun accord actuel'],['Financement à obtenir',"='Financements'!E15",'Budget moins accords monétaires saisis'],['Occupants prévus',"='Paramètres'!B10+'Paramètres'!B12",'Participants + équipe maximale ; vérifier la capacité']])
fmt(kpi,[2],PCT,11,11);fmt(kpi,[2],PCT,14,14);fmt(kpi,[2],EUR,15,18);kpi['B19'].number_format=PCT;fmt(kpi,[2],EUR,20,23)
bilan=sheet('Bilan','18 / RETOUR D’EXPÉRIENCE','Trame à compléter après réalisation et clôture. Aucun bilan d’événement n’est prérempli.', ['Axe','Fait observé / preuve','Écart / enseignement','Action pour la suite','Pilote','Échéance'],[27,69,62,62,27,20],[[axis,None,None,None,None,None] for axis in ['Mobilisation du public','Accueil et accessibilité','Programme et animation','Technique et logistique','Communication et relais','Budget et financement','Équipe et coordination','Suites éditoriales']])
inputs(bilan,[2,3,4,5,6]);fmt(bilan,[6],DATE)
guide=sheet('Mode emploi','19 / MODE D’EMPLOI','Parcours recommandé : paramètres → budget → tâches → validations → déroulé → résultats.', ['Sujet','Instruction','Onglet'],[31,110,28],[
 ['Statut','Projet personnel de conception pour avril 2027. Le 8 avril est une référence non confirmée.','Paramètres'],['Saisie','Bleu = saisie, vert = formule liée ; modifier uniquement les valeurs d’entrée.','Tous'],['Budget','20 lignes détaillées, total initial 2 310 € TTC, réserve 10 %. Saisir les devis sans les confondre avec les engagements.','Budget'],['Financement','Différencier intention, accord écrit et encaissement. Les apports en nature sont hors trésorerie.','Financements'],['Dépenses','Registre vierge, une pièce par ligne, ID budget obligatoire. Utiliser 0 pour un montant réel nul.','Dépenses'],['Planning','31 tâches à dates relatives, statut et avancement. La date de revue est saisie pour éviter des alertes périmées.','Rétroplanning'],['Gantt','Vue à 10 semaines liée au planning. Les couleurs reflètent les statuts, aucune date confirmée implicitement.','Gantt'],['Dépendances','Contrôle du prédécesseur principal ; les autres conditions d’ouverture figurent dans Décisions.','Décisions'],['Résultats','Inscriptions, entrées et questionnaires sont agrégés et vierges. Aucun contact personnel n’est inclus.','Indicateurs'],['Capacité','Comparer les 88 occupants prévus (80 public + 8 équipe) à la capacité vérifiée du lieu.','Paramètres'],['RACI','Une personne arbitre chaque action ; attribuer les rôles proposés après constitution de l’équipe.','RACI'],['Mise à jour','Après changement de jauge/budget, adapter aussi les PDF de conception et les communications avant diffusion.','Budget'],['Version','V2 · 13 octobre 2026 · Farès Benamar. Documents sources et visuels conservés dans le dossier.','Mode emploi']])
for r in range(8,21):guide.cell(r,3).hyperlink=f"#'{guide.cell(r,3).value}'!A1" if guide.cell(r,3).value in wb.sheetnames else "#'Tableau de bord'!A1";guide.cell(r,3).font=Font(color=BLUE,underline='single')

# Dashboard: formulas, finance category chart and a task status chart.
d=wb['Tableau de bord'];d._tables.clear();d.auto_filter.ref=None;d.freeze_panes='A8';d.column_dimensions['D'].width=3
metrics=[('Budget prévisionnel TTC',"='Budget'!H36",'20 lignes de coût + réserve 10 %',EUR),('Marge sur plafond',"='Budget'!H38",'Plafond de scénario : 2 500 €',EUR),('Financement confirmé',"='Financements'!E14",'Aucun accord saisi actuellement',EUR),('Reste à financer',"='Financements'!E15",'Budget moins ressources confirmées',EUR),('Participants visés',"='Paramètres'!B10",'Objectif ; pas une fréquentation acquise','0'),('Occupants prévus',"='Indicateurs'!B24",'Public + équipe ; capacité à vérifier','0'),('Tâches terminées','=COUNTIF(\'Rétroplanning\'!K8:K38,"Terminé")','Sur 31 tâches du plan','0'),('Avancement tâches',"='Indicateurs'!B19",'Proportion des tâches terminées',PCT),('Tâches en retard','=COUNTIF(\'Rétroplanning\'!N8:N38,"En retard")','Calcul depuis la date de revue','0'),('Validations non acquises','=COUNTIF(\'Décisions\'!F8:F15,"<>Validé")','Préconditions de diffusion et ouverture','0'),('Risques prioritaires','=COUNTIF(\'Risques\'!E8:E15,">=6")','Probabilité × impact ≥ 6','0'),('Paiements saisis',"='Indicateurs'!B17",'Vide tant que le registre est vierge',EUR)]
for r,(label,value,note,nf) in enumerate(metrics,8):
 d.cell(r,1,label);d.cell(r,2,value);d.cell(r,3,note);d.row_dimensions[r].height=31
 for c in d[r][:3]:c.fill=PatternFill('solid',fgColor=PAPER if r%2==0 else 'FFFFFF');c.font=Font(name='Aptos',size=11,color=INK);c.alignment=Alignment(vertical='center',wrap_text=True)
 d.cell(r,2).alignment=Alignment(horizontal='right',vertical='center',indent=1);d.cell(r,3).alignment=Alignment(wrap_text=True,vertical='center',indent=1);d.cell(r,2).font=Font(name='Aptos',size=18,bold=True,color=GREEN);d.cell(r,2).number_format=nf
cats=['Lieu','Technique','Intervenants','Collation','Supports','Administration','Accessibilité','Documentation']
d['C13']="=IF(B13>'Paramètres'!B11,\"Capacité minimale à augmenter\",\"Public + équipe ; capacité à vérifier\")"
d['A23']='RÉPARTITION DU BUDGET DE BASE';d['A23'].font=Font(bold=True,color=GOLD)
for r,cat in enumerate(cats,24):d.cell(r,1,cat);d.cell(r,2,f'=SUMIF(Budget!B8:B27,A{r},Budget!H8:H27)');d.cell(r,2).number_format=EUR
d['A34']='ÉTAT DES TÂCHES';d['A34'].font=Font(bold=True,color=GOLD)
for r,stat in enumerate(['À faire','En cours','Terminé','Bloqué'],35):d.cell(r,1,stat);d.cell(r,2,f'=COUNTIF(\'Rétroplanning\'!K8:K38,A{r})')
chart=BarChart();chart.title='Budget par enveloppe · TTC';chart.add_data(Reference(d,min_col=2,min_row=24,max_row=31));chart.set_categories(Reference(d,min_col=1,min_row=24,max_row=31));chart.type='bar';chart.legend=None;chart.series[0].graphicalProperties.solidFill=GOLD;chart.series[0].graphicalProperties.line.solidFill=GOLD;chart.height=10;chart.width=19;chart.x_axis.title='€ TTC';d.add_chart(chart,'E7')
chart=DoughnutChart();chart.title='Plan d’action · statuts';chart.add_data(Reference(d,min_col=2,min_row=35,max_row=38));chart.set_categories(Reference(d,min_col=1,min_row=35,max_row=38));chart.holeSize=65;chart.series[0].data_points=[DataPoint(idx=i,spPr=__import__('openpyxl').chart.shapes.GraphicalProperties(solidFill=c)) for i,c in enumerate([GOLD,BLUE,GREEN,RED])];chart.height=9;chart.width=19;d.add_chart(chart,'E27')
for j in range(5,16):d.column_dimensions[L(j)].width=12
d['A41']='ACCÈS RAPIDE';d['A41'].font=Font(bold=True,color=GOLD)
for r,name in enumerate(wb.sheetnames[1:],42):d.cell(r,1,name);d.cell(r,1).hyperlink=f"#'{name}'!A1";d.cell(r,1).font=Font(color=BLUE,underline='single');d.row_dimensions[r].height=22
d['A63']='Légende : bleu = saisie · vert = formule · projet non réalisé';d.print_area='A1:O63';d.page_setup.paperSize=d.PAPERSIZE_A3;d.page_setup.fitToHeight=2

# Central validation ranges.
lists=wb.create_sheet('Listes');lists.sheet_state='hidden'
for r,v in enumerate(['À faire','En cours','Terminé','Bloqué'],1):lists.cell(r,1,v)
for name,ref in [('Statuts',"'Listes'!$A$1:$A$4"),('BudgetIDs',"'Budget'!$A$8:$A$27"),('TaskIDs',"'Rétroplanning'!$A$8:$A$38")]:wb.defined_names.add(DefinedName(name,attr_text=ref))
for w in wb:
 if w.title=='Listes':continue
 for row in w:
  for c in row:
   if c.data_type=='f':c.comment=Comment('Formule de pilotage. Les données réelles restent vides avant saisie.','Farès Benamar')
wb.active=0
wb.save(OUT/'07-classeur-pilotage.xlsx')
print('Built workbook:',len(wb.sheetnames)-1,'visible sheets, 31 tasks, 20 cost lines; 2 310 EUR baseline.')
